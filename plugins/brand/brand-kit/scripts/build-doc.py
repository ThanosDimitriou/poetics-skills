#!/usr/bin/env python3
"""build-doc.py - render a Poetics document as an editable .docx (python-docx).

Runs standalone, no AI required:
    python3 build-doc.py <content.json> <output.docx>

Same content.json as build-doc-pdf.py (schemas/doc-content-schema.md). The PDF is the version to
send; the .docx is for people who need to edit in Word.

Documents follow brand-kit/design-system.md (the website's design system): A4, off-white paper
(Word shows it as the page colour; it doesn't print unless "print background colours" is on),
ink text, Instrument Serif headings, Inter Tight text, hairlines instead of filled boxes, the
website's wordmark, "poeticsvr.com · Confidential" and the page number in the footer.
"theme": "dark" switches to the dark set.

Fonts are named, not embedded (Word needs them installed): install
assets/fonts/InstrumentSerif-Regular.ttf and InterTight-Variable.ttf, or Word falls back to
Georgia and Arial. Inter Tight's 500 weight isn't selectable in Word, so headings use size, not
weight. There are no italics: <em> renders upright.
"""

import json
import pathlib
import sys
from html.parser import HTMLParser


# -- Preflight checks ---------------------------------------------------------

def check_dependencies():
    """Ensure python-docx is importable; install it on a fresh sandbox if not."""
    import importlib
    import subprocess

    try:
        importlib.import_module("docx")
        return
    except ImportError:
        pass

    result = subprocess.run(
        [sys.executable, "-m", "pip", "install", "python-docx", "--break-system-packages", "--quiet"],
    )
    if result.returncode != 0:
        # Older pip (pre-PEP 668) doesn't recognize --break-system-packages.
        subprocess.run([sys.executable, "-m", "pip", "install", "python-docx", "--quiet"], check=True)
    importlib.invalidate_caches()
    try:
        importlib.import_module("docx")
    except ImportError:
        sys.exit("error: python-docx is required but could not be installed")


check_dependencies()

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Emu, Inches, Pt, RGBColor  # noqa: E402


# -- Spec constants (from brand.md) --------------------------------------------

BODY_FONT = "Inter Tight"
DISPLAY_FONT = "Instrument Serif"
FALLBACK_FONT = "Arial"
DISPLAY_FALLBACK = "Georgia"
LINE = 1.35

PAGE_W_IN = 210 / 25.4          # A4
PAGE_H_IN = 297 / 25.4
MARGIN_SIDE_IN = 20 / 25.4
MARGIN_TOP_IN = 30 / 25.4
MARGIN_BOTTOM_IN = 22 / 25.4

# design-system.md colour roles, light theme. main() swaps in the dark set for "theme": "dark".
THEMES = {
    "light": {"paper": "EFF0EC", "paper_deep": "E5E8E3", "shade": "C5CBC4", "ink": "121614",
              "ink_soft": "5B625C", "lens": "16249C", "wordmark": "wordmark-mono-ink.png"},
    "dark": {"paper": "121614", "paper_deep": "1C211E", "shade": "3A423D", "ink": "F4F5F3",
             "ink_soft": "8A918B", "lens": "8F9CFF", "wordmark": "wordmark-mono-paper.png"},
}
PAPER, PAPER_DEEP, SHADE, INK, INK_SOFT, LENS = (THEMES["light"][k] for k in
                                                ("paper", "paper_deep", "shade", "ink", "ink_soft", "lens"))

# Type scale: (size_pt, bold, colour role, display_font). Colours resolve at render time.
STYLES = {
    "title":    (44, False, "ink", True),
    "subtitle": (13, False, "ink_soft", False),
    "date":     (9, False, "ink_soft", False),
    "h1":       (26, False, "ink", True),
    "h2":       (13.5, False, "ink", False),
    "h3":       (11.5, False, "ink", False),
    "h4":       (10.5, False, "ink", False),
    "h5":       (9, False, "ink_soft", False),
    "h6":       (9, False, "ink_soft", False),
    "body":     (10.5, False, "ink", False),
}


def C(role):
    """A colour role's hex for the current theme."""
    return globals()[role.upper()]

HEADING_SPACE = {
    "h1": (0, 14), "h2": (18, 5), "h3": (14, 4),
    "h4": (12, 4), "h5": (12, 4), "h6": (12, 4),
}

VOID_TAGS = {"br", "img", "hr"}
BLOCK_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "ul", "ol", "table", "div", "hr", "figure"}

KIT_ROOT = pathlib.Path(__file__).resolve().parent.parent
WORDMARK = KIT_ROOT / "assets/logos" / THEMES["light"]["wordmark"]


def set_page_colour(doc, hex_fill):
    """Word's page colour: shown on screen, printed only with 'print background colours' on."""
    bg = OxmlElement("w:background")
    bg.set(qn("w:color"), hex_fill)
    doc.element.insert(0, bg)
    settings = doc.settings.element
    if settings.find(qn("w:displayBackgroundShape")) is None:
        settings.append(OxmlElement("w:displayBackgroundShape"))


def rule(element_pr_parent, side, hex_color, sz=4, space=8):
    """A hairline on one side of a paragraph."""
    pPr = element_pr_parent.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), hex_color)
    pBdr.append(b)


# -- Low-level XML helpers ----------------------------------------------------

def _shade(element_pr, hex_fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    element_pr.append(shd)


def set_cell_bg(cell, hex_fill):
    _shade(cell._tc.get_or_add_tcPr(), hex_fill)


def set_table_borders(table, hex_color, sz=6):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), hex_color)
        borders.append(e)
    tblPr.append(borders)


def set_table_cell_margins(table, top, left=None, bottom=None, right=None):
    if left is None:
        left = bottom = right = top
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), str(int(val * 20)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def pad_cell_paragraphs(cell, h_pt):
    for p in cell.paragraphs:
        p.paragraph_format.left_indent = Pt(h_pt)
        p.paragraph_format.right_indent = Pt(h_pt)


CONTENT_W_IN = PAGE_W_IN - 2 * MARGIN_SIDE_IN


def align_table_full_width(table):
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.insert(0, tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(CONTENT_W_IN * 1440)))

    for old in tblPr.findall(qn("w:tblInd")):
        tblPr.remove(old)
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    jc = tblPr.find(qn("w:jc"))
    if jc is not None:
        jc.addnext(ind)
    else:
        tblW.addnext(ind)

    table.autofit = False
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_col_widths(table, widths_in):
    for i, w in enumerate(widths_in):
        table.columns[i].width = Inches(w)
        for cell in table.columns[i].cells:
            cell.width = Inches(w)


def set_theme_fonts(doc, typeface):
    import re
    theme_rt = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
    for rel in doc.part.rels.values():
        if rel.reltype != theme_rt:
            continue
        tp = rel.target_part
        try:
            blob = tp.blob.decode("utf-8")
        except Exception:
            return

        def _swap(block_tag, xml):
            return re.sub(
                rf'(<a:{block_tag}>\s*<a:latin typeface=")[^"]*(")',
                rf"\g<1>{typeface}\g<2>",
                xml,
                count=1,
                flags=re.S,
            )

        blob = _swap("majorFont", blob)
        blob = _swap("minorFont", blob)
        try:
            tp._blob = blob.encode("utf-8")
        except Exception:
            pass
        return


def add_font_fallback(doc, primary, alt=FALLBACK_FONT):
    font_rt = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/fontTable"
    for rel in doc.part.rels.values():
        if rel.reltype != font_rt:
            continue
        tp = rel.target_part
        try:
            blob = tp.blob.decode("utf-8")
        except Exception:
            return
        if f'w:name="{primary}"' in blob or "</w:fonts>" not in blob:
            return
        entry = (
            f'<w:font w:name="{primary}">'
            f'<w:altName w:val="{alt}"/>'
            f'<w:panose1 w:val="020B0604020202020204"/>'
            f'<w:charset w:val="00"/>'
            f'<w:family w:val="swiss"/>'
            f'<w:pitch w:val="variable"/>'
            f"</w:font>"
        )
        blob = blob.replace("</w:fonts>", entry + "</w:fonts>")
        try:
            tp._blob = blob.encode("utf-8")
        except Exception:
            pass
        return


def add_page_number(paragraph, label="poeticsvr.com · Confidential"):
    """Footer: the site and the confidential mark at the left, the page number at the right."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(PAGE_W_IN - 2 * MARGIN_SIDE_IN),
                                                      WD_TAB_ALIGNMENT.RIGHT)
    style_run(paragraph.add_run(label + "\t"), 8, False, C("ink_soft"))
    run = paragraph.add_run()
    style_run(run, 8, False, C("ink_soft"))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_hyperlink(paragraph, url, runs_fmt_text):
    """Links use the lens colour: blue means you can act here (design-system.md)."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    for text, fmt in runs_fmt_text:
        font = BODY_FONT
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        rFonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rFonts.set(qn(attr), font)
        rPr.append(rFonts)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(STYLES["body"][0] * 2)))
        rPr.append(sz)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), C("lens"))
        rPr.append(color)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
        if fmt.get("bold"):
            rPr.append(OxmlElement("w:b"))
        r.append(rPr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        hyperlink.append(r)
    paragraph._p.append(hyperlink)


# -- Run / paragraph styling --------------------------------------------------

def style_run(run, size_pt, bold, hex_color, italic=False, display=False):
    """Display -> Instrument Serif (headings, figures), else Inter Tight. hex_color may be a
    colour role ("ink", "ink_soft", ...). There are no italics in the system: always upright."""
    font = DISPLAY_FONT if display else BODY_FONT
    if hex_color in THEMES["light"]:
        hex_color = C(hex_color)
    italic = False
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), font)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(hex_color)


def new_paragraph(container, space_before=None, space_after=None, line=LINE):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line
    if space_before is not None:
        pf.space_before = Pt(space_before)
    if space_after is not None:
        pf.space_after = Pt(space_after)
    return p


HEADING_STYLE_NAMES = {
    "h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3",
    "h4": "Heading 4", "h5": "Heading 5", "h6": "Heading 6",
}


def add_heading(container, level, text):
    size, bold, color, display = STYLES[level]
    before, after = HEADING_SPACE[level]
    p = new_paragraph(container, space_before=before, space_after=after)
    style_name = HEADING_STYLE_NAMES.get(level)
    if style_name:
        try:
            p.style = style_name
        except KeyError:
            pass
    p.paragraph_format.keep_with_next = True
    style_run(p.add_run(text), size, bold, color, display=display)
    return p


# -- HTML subset parsing ------------------------------------------------------

class TreeBuilder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = {"tag": "root", "attrs": {}, "children": []}
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = {"tag": tag, "attrs": dict(attrs), "children": []}
        self.stack[-1]["children"].append(node)
        if tag not in VOID_TAGS:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1]["children"].append({"tag": tag, "attrs": dict(attrs), "children": []})

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i]["tag"] == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        self.stack[-1]["children"].append(data)


def parse_html(raw):
    tb = TreeBuilder()
    tb.feed(raw or "")
    return tb.root["children"]


def has_class(node, name):
    return name in (node.get("attrs", {}).get("class", "").split())


# -- Inline rendering ---------------------------------------------------------

def collect_inline_segments(children, fmt):
    segs = []
    for ch in children:
        if isinstance(ch, str):
            text = ch.replace("\n", " ")
            if text:
                segs.append(("text", text, dict(fmt)))
            continue
        tag = ch["tag"]
        if tag == "br":
            segs.append(("break", None, None))
        elif tag in ("strong", "b"):
            segs.extend(collect_inline_segments(ch["children"], {**fmt, "bold": True}))
        elif tag in ("em", "i"):
            segs.extend(collect_inline_segments(ch["children"], {**fmt, "italic": True}))
        elif tag == "a":
            url = ch["attrs"].get("href", "")
            inner = collect_inline_segments(ch["children"], fmt)
            link_runs = [(t, f) for kind, t, f in inner if kind == "text"]
            if url and link_runs:
                segs.append(("link", url, link_runs))
            else:
                segs.extend(inner)
        else:
            segs.extend(collect_inline_segments(ch["children"], fmt))
    return segs


def render_inline(paragraph, children, base_size, base_color):
    for kind, val, fmt in collect_inline_segments(children, {}):
        if kind == "break":
            paragraph.add_run().add_break()
        elif kind == "text":
            run = paragraph.add_run(val)
            style_run(run, base_size, fmt.get("bold", False), base_color, fmt.get("italic", False))
        elif kind == "link":
            add_hyperlink(paragraph, val, fmt)


# -- Block rendering ----------------------------------------------------------

def render_blocks(container, nodes):
    for node in nodes:
        if isinstance(node, str):
            if node.strip():
                p = new_paragraph(container, space_after=12)
                style_run(p.add_run(node.strip()), STYLES["body"][0], False, STYLES["body"][2])
            continue
        render_block(container, node)


def render_block(container, node):
    tag = node["tag"]
    if tag in ("h2", "h3", "h4", "h5", "h6"):
        add_heading(container, tag, inline_text(node["children"]))
    elif tag == "p":
        p = new_paragraph(container, space_after=12)
        render_inline(p, node["children"], STYLES["body"][0], STYLES["body"][2])
    elif tag in ("ul", "ol"):
        style_name = "List Number" if tag == "ol" else "List Bullet"
        for li in node["children"]:
            if isinstance(li, dict) and li["tag"] == "li":
                p = container.add_paragraph(style=style_name)
                p.paragraph_format.line_spacing = LINE
                p.paragraph_format.space_after = Pt(6)
                render_inline(p, li["children"], STYLES["body"][0], STYLES["body"][2])
    elif tag == "img":
        render_image(container, node)
    elif tag == "figure":
        render_figure(container, node)
    elif tag == "hr":
        render_hr(container)
    elif tag == "table":
        render_table(container, node)
    elif tag == "div":
        if has_class(node, "callout-2col"):
            render_callout_2col(container, node)
        elif has_class(node, "price-callout"):
            render_price_callout(container, node["children"])
        elif has_class(node, "callout"):
            render_callout(container, node["children"])
        elif has_class(node, "pull-quote"):
            render_pull_quote(container, node["children"])
        elif has_class(node, "stat-row"):
            render_stat_row(container, node["children"])
        elif has_class(node, "eyebrow"):
            render_eyebrow(container, node["children"])
        else:
            render_blocks(container, node["children"])


IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".webp"}
CONTENT_WIDTH = Inches(6.5)


def render_image(container, node):
    src = node.get("attrs", {}).get("src", "")
    if not src or src.startswith(("http", "data:")):
        return
    path = pathlib.Path(src)
    if path.suffix.lower() not in IMAGE_EXTS or not path.exists():
        return
    p = new_paragraph(container, space_after=12)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    try:
        from docx.image.exceptions import UnrecognizedImageError
        try:
            p.add_run().add_picture(str(path), width=CONTENT_WIDTH)
        except UnrecognizedImageError:
            return
    except Exception:
        return


def render_figure(container, node):
    """<figure><img src="..."><figcaption>...</figcaption></figure>"""
    img_node = None
    caption = None
    for c in node["children"]:
        if not isinstance(c, dict):
            continue
        if c["tag"] == "img":
            img_node = c
        elif c["tag"] == "figcaption":
            caption = inline_text(c["children"])
    if img_node:
        render_image(container, img_node)
    if caption:
        p = new_paragraph(container, space_after=12)
        style_run(p.add_run(caption), 8.5, False, C("ink_soft"))


def render_hr(container):
    p = new_paragraph(container, space_before=16, space_after=16)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), C("shade"))
    pBdr.append(bottom)
    pPr.append(pBdr)


def render_eyebrow(container, children):
    """An overline: caption size, ink-soft, sentence case, no tracking."""
    p = new_paragraph(container, space_before=0, space_after=6)
    style_run(p.add_run(inline_text(children)), 9, False, C("ink_soft"))


def render_pull_quote(container, children):
    text_node, attribution = None, ""
    for ch in children:
        if not isinstance(ch, dict):
            continue
        if ch["tag"] == "div" and has_class(ch, "pull-quote-text"):
            text_node = ch
        elif ch["tag"] == "div" and has_class(ch, "pull-quote-attribution"):
            attribution = inline_text(ch["children"])
    if text_node is None:
        return
    p = new_paragraph(container, space_before=12, space_after=4 if attribution else 12, line=1.4)
    p.paragraph_format.left_indent = Pt(20)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "4")
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), C("shade"))
    pBdr.append(left)
    pPr.append(pBdr)
    render_inline(p, text_node["children"], 14, C("ink"))
    if attribution:
        ap = new_paragraph(container, space_after=14)
        ap.paragraph_format.left_indent = Pt(20)
        style_run(ap.add_run(attribution), 9.5, False, C("ink_soft"))


def render_stat_row(container, children):
    stats = [c for c in children if isinstance(c, dict) and c["tag"] == "div" and has_class(c, "stat")]
    if not stats:
        return
    n = len(stats)
    table = container.add_table(rows=1, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    align_table_full_width(table)
    set_col_widths(table, [CONTENT_W_IN / n] * n)
    set_table_cell_margins(table, 6, 8, 4, 8)
    for i, stat in enumerate(stats):
        cell = table.cell(0, i)
        fill_stat_cell(cell, stat["children"])
    new_paragraph(container, space_after=8)


def fill_stat_cell(cell, src):
    cell.paragraphs[0].text = ""
    first = True
    for ch in src:
        if not isinstance(ch, dict):
            continue
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.line_spacing = 1.1
        if ch["tag"] == "div" and has_class(ch, "stat-value"):
            p.paragraph_format.space_after = Pt(4)
            rule(p._p, "top", C("shade"), sz=4, space=8)
            style_run(p.add_run(inline_text(ch["children"])), 32, False, C("ink"), display=True)
        elif ch["tag"] == "div" and has_class(ch, "stat-label"):
            p.paragraph_format.space_after = Pt(2)
            render_inline(p, ch["children"], 9.5, C("ink_soft"))
        elif ch["tag"] == "div" and has_class(ch, "stat-sublabel"):
            render_inline(p, ch["children"], 8.5, C("ink_soft"))
        else:
            first = True  # unrecognized child, don't consume a blank paragraph


def render_price_callout(container, children):
    table = container.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    align_table_full_width(table)
    set_col_widths(table, [CONTENT_W_IN])
    set_table_cell_margins(table, CALLOUT_PAD_V, 0, CALLOUT_PAD_V, 0)
    cell = table.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")
    top.set(qn("w:space"), "0")
    top.set(qn("w:color"), C("shade"))
    tcBorders.append(top)
    tcPr.append(tcBorders)
    fill_price_cell(cell, children)
    new_paragraph(container, space_after=4)


def fill_price_cell(cell, src):
    cell.paragraphs[0].text = ""
    first = True
    for ch in src:
        if not isinstance(ch, dict):
            continue
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.line_spacing = LINE
        if ch["tag"] == "div" and has_class(ch, "price-value"):
            p.paragraph_format.space_after = Pt(2)
            style_run(p.add_run(inline_text(ch["children"])), 38, False, C("ink"), display=True)
        elif ch["tag"] == "div" and has_class(ch, "price-note"):
            p.paragraph_format.space_after = Pt(6)
            style_run(p.add_run(inline_text(ch["children"])), 10, False, C("ink_soft"))
        elif ch["tag"] == "div" and has_class(ch, "price-body"):
            render_inline(p, ch["children"], 10.5, C("ink"))
        else:
            first = True


def inline_text(children):
    return "".join(t for kind, t, _ in collect_inline_segments(children, {}) if kind == "text").strip()


def iter_rows(table_node):
    rows = []

    def walk(n):
        for ch in n.get("children", []):
            if isinstance(ch, dict):
                if ch["tag"] == "tr":
                    rows.append(ch)
                else:
                    walk(ch)
    walk(table_node)
    return rows


def render_table(container, table_node):
    rows = iter_rows(table_node)
    if not rows:
        return
    matrix = []
    header_flags = []
    for tr in rows:
        cells = [c for c in tr["children"] if isinstance(c, dict) and c["tag"] in ("td", "th")]
        matrix.append(cells)
        header_flags.append(any(c["tag"] == "th" for c in cells))
    ncols = max((len(r) for r in matrix), default=0)
    if ncols == 0:
        return

    table = container.add_table(rows=len(matrix), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    align_table_full_width(table)
    set_col_widths(table, [CONTENT_W_IN / ncols] * ncols)
    set_table_borders(table, C("shade"), sz=4)
    set_table_cell_margins(table, 5)

    for ri, cells in enumerate(matrix):
        is_header = header_flags[ri]
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            node = cells[ci] if ci < len(cells) else None
            recommended = node is not None and has_class(node, "recommended")
            if recommended:
                set_cell_bg(cell, C("paper_deep"))
            render_cell(cell, node, bold=False, header=is_header, first=(ci == 0 and not is_header))
    new_paragraph(container, space_after=4)


def render_cell(cell, node, bold=False, header=False, first=False):
    cell.paragraphs[0].text = ""
    if node is None:
        return
    block_children = [c for c in node["children"] if isinstance(c, dict) and c["tag"] in BLOCK_TAGS]
    if block_children:
        for c in node["children"]:
            if isinstance(c, dict) and c["tag"] in BLOCK_TAGS:
                render_block(cell, c)
    else:
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = LINE
        size, color = (8.5, C("ink_soft")) if header else (9.5, C("ink"))
        segs = collect_inline_segments(node["children"], {"bold": bold})
        if not segs:
            return
        for kind, val, fmt in segs:
            if kind == "break":
                p.add_run().add_break()
            elif kind == "text":
                style_run(p.add_run(val), size, fmt.get("bold", False), color, fmt.get("italic", False))
            elif kind == "link":
                add_hyperlink(p, val, fmt)


CALLOUT_PAD_V = 8


def render_callout(container, children):
    table = container.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    align_table_full_width(table)
    set_col_widths(table, [CONTENT_W_IN])
    set_table_cell_margins(table, CALLOUT_PAD_V, 0, CALLOUT_PAD_V, 0)
    cell = table.cell(0, 0)
    fill_callout_cell(cell, children)
    new_paragraph(container, space_after=4)
    return table


def fill_callout_cell(cell, src):
    """A hairline above, title in ink, body in ink-soft. Never a filled box."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")
    top.set(qn("w:space"), "0")
    top.set(qn("w:color"), C("shade"))
    tcBorders.append(top)
    tcPr.append(tcBorders)
    cell.paragraphs[0].text = ""
    first_done = False
    for ch in src:
        if isinstance(ch, dict) and ch["tag"] == "div" and has_class(ch, "callout-title"):
            p = cell.paragraphs[0] if not first_done else cell.add_paragraph()
            p.paragraph_format.line_spacing = LINE
            p.paragraph_format.space_after = Pt(6)
            style_run(p.add_run(inline_text(ch["children"])), 11, False, C("ink"))
            first_done = True
        elif isinstance(ch, dict) and ch["tag"] == "div" and has_class(ch, "callout-body"):
            p = cell.paragraphs[0] if not first_done else cell.add_paragraph()
            p.paragraph_format.line_spacing = LINE
            render_inline(p, ch["children"], 10.5, C("ink_soft"))
            first_done = True
        elif isinstance(ch, dict) and ch["tag"] in BLOCK_TAGS:
            render_block(cell, ch)
            first_done = True
        elif isinstance(ch, str) and ch.strip():
            p = cell.paragraphs[0] if not first_done else cell.add_paragraph()
            p.paragraph_format.line_spacing = LINE
            style_run(p.add_run(ch.strip()), 10.5, False, C("ink_soft"))
            first_done = True


def render_callout_2col(container, node):
    callouts = [c for c in node["children"] if isinstance(c, dict) and c["tag"] == "div" and has_class(c, "callout")]
    if not callouts:
        render_blocks(container, node["children"])
        return
    gap_in = 0.2
    box_in = (CONTENT_W_IN - gap_in) / 2
    table = container.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    align_table_full_width(table)
    set_col_widths(table, [box_in, gap_in, box_in])
    set_table_cell_margins(table, CALLOUT_PAD_V, 0, CALLOUT_PAD_V, 0)
    fill_callout_cell(table.cell(0, 0), callouts[0]["children"])

    if len(callouts) > 1:
        fill_callout_cell(table.cell(0, 2), callouts[1]["children"])
        pass
    table.cell(0, 1).paragraphs[0].text = ""
    new_paragraph(container, space_after=4)


# -- Page builders ------------------------------------------------------------

def page_setup(sec, top=None):
    sec.page_width = Inches(PAGE_W_IN)
    sec.page_height = Inches(PAGE_H_IN)
    sec.top_margin = Inches(top if top is not None else MARGIN_TOP_IN)
    sec.bottom_margin = Inches(MARGIN_BOTTOM_IN)
    sec.left_margin = Inches(MARGIN_SIDE_IN)
    sec.right_margin = Inches(MARGIN_SIDE_IN)


def spacer(doc, pt):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(pt)
    return p


def build_cover(doc, title, subtitle, date, footer_label, cover_image=""):
    sec = doc.sections[0]
    page_setup(sec, top=20 / 25.4)
    footer = sec.footer
    style_run(footer.paragraphs[0].add_run(footer_label), 8, False, C("ink_soft"))

    logo_p = doc.add_paragraph()
    logo_p.paragraph_format.space_after = Pt(0)
    if WORDMARK.exists():
        logo_p.add_run().add_picture(str(WORDMARK), height=Pt(16))

    if cover_image and pathlib.Path(cover_image).exists():
        spacer(doc, 40)
        img_p = doc.add_paragraph()
        img_p.add_run().add_picture(str(cover_image), width=Inches(PAGE_W_IN - 2 * MARGIN_SIDE_IN))
        spacer(doc, 120)
    else:
        spacer(doc, 420)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(10)
    title_p.paragraph_format.line_spacing = 1.05
    style_run(title_p.add_run(title), *STYLES["title"])
    if subtitle:
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_after = Pt(8)
        style_run(sub_p.add_run(subtitle), *STYLES["subtitle"])
    if date:
        date_p = doc.add_paragraph()
        date_p.paragraph_format.space_before = Pt(36)
        style_run(date_p.add_run(date), *STYLES["date"])


def start_body_section(doc, footer_label):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(sec)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "2")
    sec._sectPr.append(pg)

    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(28)   # clear space below the wordmark (the guardrail)
    if WORDMARK.exists():
        hp.add_run().add_picture(str(WORDMARK), height=Pt(11))

    footer = sec.footer
    footer.is_linked_to_previous = False
    add_page_number(footer.paragraphs[0], footer_label)
    return sec


def add_section_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def build_separator(doc, title):
    add_section_page_break(doc)
    spacer(doc, 520)
    p = doc.add_paragraph()
    style_run(p.add_run(title), 40, False, C("ink"), display=True)


# -- Main ---------------------------------------------------------------------

def main():
    if len(sys.argv) != 3:
        sys.exit("usage: build-doc.py <content.json> <output.docx>")
    content_path, out_path = sys.argv[1], sys.argv[2]
    with open(content_path, encoding="utf-8") as f:
        data = json.load(f)

    title = data.get("title", "Untitled")
    subtitle = data.get("subtitle", "")
    date = data.get("date", "")
    sections = data.get("sections", []) or []
    separator_title = data.get("separator_title", "")

    global PAPER, PAPER_DEEP, SHADE, INK, INK_SOFT, LENS, WORDMARK
    theme = THEMES.get(data.get("theme", "light"), THEMES["light"])
    PAPER, PAPER_DEEP, SHADE, INK, INK_SOFT, LENS = (theme[k] for k in
                                                    ("paper", "paper_deep", "shade", "ink", "ink_soft", "lens"))
    WORDMARK = KIT_ROOT / "assets/logos" / theme["wordmark"]
    confidential = {"el": "Εμπιστευτικό"}.get(data.get("lang"), "Confidential")
    footer_label = "poeticsvr.com" + (f" · {confidential}" if data.get("confidential", True) else "")
    cover_image = data.get("cover_image", "")
    if cover_image.startswith("photo:"):
        try:
            import importlib.util as _ilu
            spec = _ilu.spec_from_file_location("photo_library", pathlib.Path(__file__).parent / "photo-library.py")
            lib = _ilu.module_from_spec(spec)
            spec.loader.exec_module(lib)
            cover_image = str(lib.resolve(cover_image[6:]) or "")
        except Exception:
            cover_image = ""
    elif cover_image and not pathlib.Path(cover_image).is_absolute():
        cover_image = str(pathlib.Path(content_path).resolve().parent / cover_image)

    doc = Document()
    set_page_colour(doc, PAPER)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(STYLES["body"][0])
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    set_theme_fonts(doc, FALLBACK_FONT)
    add_font_fallback(doc, BODY_FONT)
    add_font_fallback(doc, DISPLAY_FONT, alt=DISPLAY_FALLBACK)
    normal.font.color.rgb = RGBColor.from_string(INK)

    build_cover(doc, title, subtitle, date, footer_label, cover_image)
    start_body_section(doc, footer_label)

    for i, section in enumerate(sections):
        if i > 0:
            add_section_page_break(doc)
        heading = section.get("heading", "")
        if heading:
            add_heading(doc, "h1", heading)
        render_blocks(doc, parse_html(section.get("html", "")))

    if separator_title:
        build_separator(doc, separator_title)

    doc.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
