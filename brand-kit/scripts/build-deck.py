#!/usr/bin/env python3
"""build-deck.py - render a branded Poetics slide deck from a deck.json spec.

Runs standalone, no AI required:
    python3 build-deck.py <deck.json> <output.html|output.pdf|output.pptx>

The output format is chosen by the output file's extension:
  .html  - self-contained deck, zero external dependencies, opens in any
           browser, printable to PDF from there too. This is the primary
           deliverable — hand it to anyone, Claude or not.
  .pdf   - rendered via Playwright/Chromium from the same HTML.
  .pptx  - editable PowerPoint via python-pptx (styled but not pixel-identical
           to the HTML — python-pptx text boxes approximate the layout; treat
           the HTML/PDF as the source of visual truth, PPTX as the editable
           fallback for people who need to tweak slides in PowerPoint).

deck.json schema:
{
  "title": "Deck title",           // used for the HTML <title> and PPTX metadata
  "slides": [
    {"type": "title", "title": "...", "subtitle": "..."},
    {"type": "section", "title": "Section name"},
    {"type": "content", "title": "...", "bullets": ["...", "..."]},
    {"type": "cards", "title": "...", "cards": [{"title": "...", "body": "..."}]},
    {"type": "quote", "quote": "...", "attribution": "..."},
    {"type": "closing", "title": "...", "cta": "..."}
  ]
}

Every text field is treated as plain text and HTML-escaped automatically —
callers do not need to escape content themselves (unlike build-doc-pdf.py's
richer "html" field). Wrap 2-4 word phrases in *single asterisks* to render
them in the Cormorant Garamond italic "emphasis" voice per brand.md, e.g.
"A new way *to experience* art".
"""

import base64
import html as html_lib
import json
import pathlib
import re
import sys

KIT_ROOT = pathlib.Path(__file__).resolve().parent.parent
FONTS = KIT_ROOT / "assets/fonts"
LOGOS = KIT_ROOT / "assets/logos"
BACKGROUNDS = KIT_ROOT / "assets/backgrounds"

ACCENT = "#0099FF"
CANVAS = "#06070F"          # deep blue-black, matching the Athens College deck
RAISED = "#111214"
TEXT_1 = "#F4F4F2"
TEXT_2 = "#9A9A98"


def e(text):
    return html_lib.escape(str(text or ""))


EMPH_RE = re.compile(r"\*([^*]+)\*")


def emph(text):
    """Escape text, then render *phrase* as Cormorant Garamond italic emphasis."""
    parts = []
    last = 0
    raw = str(text or "")
    for m in EMPH_RE.finditer(raw):
        parts.append(e(raw[last:m.start()]))
        parts.append(f'<em>{e(m.group(1))}</em>')
        last = m.end()
    parts.append(e(raw[last:]))
    return "".join(parts)


def b64(path: pathlib.Path) -> str:
    return base64.b64encode(path.read_bytes()).decode()


def font_face_css():
    faces = [
        ("Inter", 300, "normal", "Inter-Light.ttf"),
        ("Inter", 400, "normal", "Inter-Regular.ttf"),
        ("Inter", 500, "normal", "Inter-Medium.ttf"),
        ("Inter", 600, "normal", "Inter-SemiBold.ttf"),
        ("Inter", 700, "normal", "Inter-Bold.ttf"),
        ("Cormorant Garamond", 400, "italic", "CormorantGaramond-Italic.ttf"),
        ("Space Mono", 400, "normal", "SpaceMono-Regular.ttf"),
    ]
    out = []
    for family, weight, style, filename in faces:
        path = FONTS / filename
        if not path.exists():
            continue
        out.append(
            f"@font-face {{ font-family: '{family}'; font-weight: {weight}; "
            f"font-style: {style}; src: url('data:font/ttf;base64,{b64(path)}') format('truetype'); }}"
        )
    return "\n".join(out)


DECK_CSS = """
* { box-sizing: border-box; margin: 0; padding: 0; }
html, body { background: %(canvas)s; }
body { font-family: 'Inter', sans-serif; color: %(text1)s; }
.slide {
  position: relative;
  width: 1920px; height: 1080px;
  overflow: hidden;
  background: %(canvas)s;
  background-size: cover;
  background-position: center;
  page-break-after: always;
  display: flex;
  flex-direction: column;
}
.slide + .slide { margin-top: 0; }
.slide-inner { position: relative; z-index: 1; flex: 1; display: flex; flex-direction: column; padding: 100px 120px; }
.logo { position: absolute; top: 56px; left: 100px; z-index: 2; height: 28px; }
.eyebrow {
  font-family: 'Space Mono', monospace; font-size: 15px; letter-spacing: 0.14em;
  text-transform: uppercase; color: %(accent)s; margin-bottom: 24px;
}
h1.slide-title {
  font-family: 'Inter', sans-serif; font-weight: 600;
  font-size: 64px; line-height: 1.08; letter-spacing: -0.02em; color: %(text1)s;
  max-width: 1100px;
}
.slide-subtitle {
  font-family: 'Inter', sans-serif; font-weight: 300; font-size: 26px;
  color: %(text2)s; margin-top: 20px; max-width: 900px;
}
em { font-family: 'Cormorant Garamond', serif; font-style: italic; color: inherit; }

.slide.title-slide .slide-inner { justify-content: flex-end; padding-bottom: 140px; }
.slide.section-slide { align-items: center; justify-content: center; }
.slide.section-slide .slide-inner { align-items: center; justify-content: center; text-align: center; }
.slide.section-slide h1.slide-title { font-size: 72px; max-width: 1400px; }

.bullets { list-style: none; margin-top: 48px; max-width: 1300px; }
.bullets li {
  font-size: 28px; line-height: 1.5; color: %(text1)s; font-weight: 300;
  padding-left: 40px; position: relative; margin-bottom: 26px;
}
.bullets li::before {
  content: ''; position: absolute; left: 0; top: 14px; width: 6px; height: 6px;
  border-radius: 1px; background: %(accent)s;
}

.cards { display: grid; grid-template-columns: repeat(2, 1fr); gap: 24px; margin-top: 56px; }
.cards.cols-3 { grid-template-columns: repeat(3, 1fr); }
.card {
  background: %(raised)s; border: 1px solid rgba(255,255,255,0.08);
  border-radius: 16px; padding: 36px 32px;
}
.card-title { font-family: 'Inter', sans-serif; font-weight: 600; font-size: 22px; color: %(text1)s; margin-bottom: 12px; }
.card-body { font-size: 18px; line-height: 1.55; color: %(text2)s; font-weight: 300; }

.quote-slide .slide-inner { justify-content: center; align-items: flex-start; }
.quote-text {
  font-family: 'Cormorant Garamond', serif; font-style: italic; font-weight: 400;
  font-size: 52px; line-height: 1.4; color: %(text1)s; max-width: 1400px;
}
.quote-attribution {
  font-family: 'Space Mono', monospace; font-size: 16px; letter-spacing: 0.1em;
  text-transform: uppercase; color: %(text2)s; margin-top: 36px;
}

.closing-slide .slide-inner { align-items: center; justify-content: center; text-align: center; }
.cta {
  margin-top: 40px; display: inline-block; font-family: 'Space Mono', monospace;
  font-size: 16px; letter-spacing: 0.08em; text-transform: uppercase; color: %(accent)s;
  border: 1px solid rgba(0,153,255,0.4); border-radius: 20px; padding: 14px 32px;
}

.footer-rule {
  position: absolute; bottom: 64px; left: 120px; right: 120px;
  border-top: 0.5px solid rgba(255,255,255,0.12); padding-top: 14px;
}
.footer-rule span {
  font-family: 'Space Mono', monospace; font-size: 13px; color: %(text2)s;
}

@media print {
  html, body { -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }
  .slide { break-after: page; }
}
""" % {"canvas": CANVAS, "raised": RAISED, "text1": TEXT_1, "text2": TEXT_2, "accent": ACCENT}


_MIME = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}


def slide_background_uri(slot: str, default_name: str) -> str:
    """Photo override from assets/backgrounds/photo/<slot>.* if present, else the gradient."""
    photo_dir = BACKGROUNDS / "photo"
    for ext in (".jpg", ".jpeg", ".png", ".webp"):
        cand = photo_dir / f"{slot}{ext}"
        if cand.exists():
            return f"data:{_MIME[ext]};base64,{b64(cand)}"
    path = BACKGROUNDS / default_name
    if path.exists():
        return f"data:image/jpeg;base64,{b64(path)}"
    return ""


def render_slide(slide: dict, logo_uri: str, footer_text: str) -> str:
    stype = slide.get("type", "content")
    if stype in ("title", "section", "closing"):
        bg_uri = slide_background_uri("deck-title", "deck-bg-title.jpg")
    else:
        bg_uri = slide_background_uri("deck-default", "deck-bg-default.jpg")
    bg_style = f"background-image:url('{bg_uri}');" if bg_uri else ""

    logo_html = f'<img class="logo" src="{logo_uri}" alt="Poetics">' if logo_uri and stype != "section" else ""
    footer_html = (
        f'<div class="footer-rule"><span>{e(footer_text)}</span></div>'
        if footer_text and stype not in ("title", "section", "closing") else ""
    )

    if stype == "title":
        body = f'''
    <div class="slide-inner">
      <h1 class="slide-title">{emph(slide.get("title", ""))}</h1>
      {f'<div class="slide-subtitle">{emph(slide.get("subtitle", ""))}</div>' if slide.get("subtitle") else ""}
    </div>'''
        cls = "title-slide"
    elif stype == "section":
        body = f'''
    <div class="slide-inner">
      {f'<div class="eyebrow">{e(slide.get("eyebrow",""))}</div>' if slide.get("eyebrow") else ""}
      <h1 class="slide-title">{emph(slide.get("title", ""))}</h1>
    </div>'''
        cls = "section-slide"
    elif stype == "content":
        bullets = "".join(f'<li>{emph(b)}</li>' for b in slide.get("bullets", []))
        body = f'''
    <div class="slide-inner">
      <h1 class="slide-title">{emph(slide.get("title", ""))}</h1>
      <ul class="bullets">{bullets}</ul>
    </div>'''
        cls = "content-slide"
    elif stype == "cards":
        cards = slide.get("cards", [])
        cols = "cols-3" if len(cards) >= 3 else ""
        card_html = "".join(
            f'<div class="card"><div class="card-title">{e(c.get("title",""))}</div>'
            f'<div class="card-body">{emph(c.get("body",""))}</div></div>'
            for c in cards
        )
        body = f'''
    <div class="slide-inner">
      <h1 class="slide-title">{emph(slide.get("title", ""))}</h1>
      <div class="cards {cols}">{card_html}</div>
    </div>'''
        cls = "cards-slide"
    elif stype == "quote":
        body = f'''
    <div class="slide-inner">
      <div class="quote-text">&ldquo;{emph(slide.get("quote", ""))}&rdquo;</div>
      {f'<div class="quote-attribution">{e(slide.get("attribution",""))}</div>' if slide.get("attribution") else ""}
    </div>'''
        cls = "quote-slide"
    elif stype == "closing":
        body = f'''
    <div class="slide-inner">
      <h1 class="slide-title">{emph(slide.get("title", ""))}</h1>
      {f'<div class="cta">{e(slide.get("cta",""))}</div>' if slide.get("cta") else ""}
    </div>'''
        cls = "closing-slide"
    else:
        body = f'<div class="slide-inner"><h1 class="slide-title">{emph(slide.get("title",""))}</h1></div>'
        cls = "content-slide"

    return f'<section class="slide {cls}" style="{bg_style}">{logo_html}{body}{footer_html}</section>'


def build_html(deck: dict) -> str:
    title = e(deck.get("title", "Poetics deck"))
    footer_text = deck.get("footer", "poeticsvr.com")
    logo_path = LOGOS / "wordmark-light.png"
    logo_uri = f"data:image/png;base64,{b64(logo_path)}" if logo_path.exists() else ""

    slides_html = "\n".join(render_slide(s, logo_uri, footer_text) for s in deck.get("slides", []))

    return f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
{font_face_css()}
{DECK_CSS}
</style>
</head>
<body>
{slides_html}
</body>
</html>"""


# -- PDF export (Playwright) ---------------------------------------------------

def export_pdf(html_path: pathlib.Path, pdf_path: str):
    import asyncio

    async def _run():
        import os as _os
        from playwright.async_api import async_playwright
        async with async_playwright() as p:
            browser = await p.chromium.launch(
                channel=_os.environ.get("POETICS_BROWSER_CHANNEL") or None)
            page = await browser.new_page(viewport={"width": 1920, "height": 1080})
            await page.goto(html_path.as_uri())
            try:
                await page.evaluate("document.fonts.ready")
                await page.wait_for_timeout(400)
            except Exception:
                await page.wait_for_timeout(1200)
            await page.pdf(
                path=pdf_path,
                width="1920px", height="1080px",
                print_background=True,
                margin={"top": "0", "bottom": "0", "left": "0", "right": "0"},
            )
            await browser.close()

    asyncio.run(_run())


def check_playwright():
    import importlib
    import subprocess
    import os as _os
    try:
        importlib.import_module("playwright")
    except ImportError:
        r = subprocess.run([sys.executable, "-m", "pip", "install", "--break-system-packages", "-q", "playwright"])
        if r.returncode != 0:
            subprocess.run([sys.executable, "-m", "pip", "install", "-q", "playwright"], check=True)
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            exe = pw.chromium.executable_path
        if not (exe and _os.path.exists(exe)):
            raise RuntimeError
    except Exception:
        import importlib.util as _ilu
        _ba_path = pathlib.Path(__file__).parent / "_brand_assets.py"
        _ba_spec = _ilu.spec_from_file_location("_brand_assets", _ba_path)
        _ba = _ilu.module_from_spec(_ba_spec)
        _ba_spec.loader.exec_module(_ba)
        channel = _ba.pick_browser_channel()
        if channel:
            # System Chrome/Edge renders identically — skip the 110 MB download
            # (which also fails outright on networks that block the Playwright CDN).
            _os.environ.setdefault("POETICS_BROWSER_CHANNEL", channel)
            print(f"Bundled Chromium not installed — using system browser (channel: {channel}).",
                  file=sys.stderr)
        else:
            subprocess.run([sys.executable, "-m", "playwright", "install", "chromium"], check=True)


# -- PPTX export (python-pptx) -------------------------------------------------

def export_pptx(deck: dict, out_path: str):
    import subprocess
    import importlib
    try:
        importlib.import_module("pptx")
    except ImportError:
        r = subprocess.run([sys.executable, "-m", "pip", "install", "python-pptx", "--break-system-packages", "--quiet"])
        if r.returncode != 0:
            subprocess.run([sys.executable, "-m", "pip", "install", "python-pptx", "--quiet"], check=True)

    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    def rgb(hexstr):
        hexstr = hexstr.lstrip("#")
        return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))

    prs = Presentation()
    prs.slide_width = Emu(12192000)   # 13.333in, 16:9
    prs.slide_height = Emu(6858000)   # 7.5in
    blank = prs.slide_layouts[6]

    def strip_emphasis(text):
        return EMPH_RE.sub(r"\1", text or "")

    for s in deck.get("slides", []):
        slide = prs.slides.add_slide(blank)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = rgb(CANVAS)

        stype = s.get("type", "content")
        title_text = strip_emphasis(s.get("title", "")) or strip_emphasis(s.get("quote", ""))

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(1.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Inter"
        p.font.size = Pt(40 if stype != "section" else 48)
        p.font.bold = True
        p.font.color.rgb = rgb(TEXT_1)

        if stype == "content" and s.get("bullets"):
            body = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.5), Inches(4.5))
            btf = body.text_frame
            btf.word_wrap = True
            for i, b in enumerate(s["bullets"]):
                bp = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
                bp.text = f"—  {strip_emphasis(b)}"
                bp.font.name = "Inter"
                bp.font.size = Pt(18)
                bp.font.color.rgb = rgb(TEXT_1)
                bp.space_after = Pt(10)
        elif stype == "quote":
            title_box.text_frame.paragraphs[0].font.name = "Cormorant Garamond"
            title_box.text_frame.paragraphs[0].font.italic = True
            if s.get("attribution"):
                attr = slide.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(8), Inches(0.6))
                ap = attr.text_frame.paragraphs[0]
                ap.text = s["attribution"]
                ap.font.name = "Inter"
                ap.font.size = Pt(13)
                ap.font.color.rgb = rgb(TEXT_2)
        elif stype == "cards" and s.get("cards"):
            cards = s["cards"]
            n = len(cards)
            col_w = Inches(11.5 / max(n, 1))
            for i, c in enumerate(cards):
                left = Inches(0.8) + col_w * i
                card_box = slide.shapes.add_textbox(left, Inches(2.6), col_w - Inches(0.2), Inches(3.5))
                ctf = card_box.text_frame
                ctf.word_wrap = True
                cp = ctf.paragraphs[0]
                cp.text = strip_emphasis(c.get("title", ""))
                cp.font.name = "Inter"
                cp.font.bold = True
                cp.font.size = Pt(16)
                cp.font.color.rgb = rgb(TEXT_1)
                bp = ctf.add_paragraph()
                bp.text = strip_emphasis(c.get("body", ""))
                bp.font.name = "Inter"
                bp.font.size = Pt(13)
                bp.font.color.rgb = rgb(TEXT_2)
        elif stype == "closing" and s.get("cta"):
            cta_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(6), Inches(0.8))
            cp = cta_box.text_frame.paragraphs[0]
            cp.text = s["cta"]
            cp.font.name = "Space Mono"
            cp.font.size = Pt(14)
            cp.font.color.rgb = rgb(ACCENT)

    prs.save(out_path)


# -- Entry point -----------------------------------------------------------------

def main():
    if len(sys.argv) != 3:
        sys.exit("usage: build-deck.py <deck.json> <output.html|output.pdf|output.pptx>")
    deck_path, out_path = sys.argv[1], sys.argv[2]
    with open(deck_path, encoding="utf-8") as f:
        deck = json.load(f)

    out = pathlib.Path(out_path)
    ext = out.suffix.lower()

    html = build_html(deck)

    if ext == ".html":
        out.write_text(html, encoding="utf-8")
        print(f"Wrote {out_path}")
    elif ext == ".pdf":
        check_playwright()
        tmp_html = pathlib.Path("/tmp") / f"{out.stem}.tmp.html"
        tmp_html.write_text(html, encoding="utf-8")
        try:
            export_pdf(tmp_html.resolve(), out_path)
        finally:
            tmp_html.unlink(missing_ok=True)
        print(f"Wrote {out_path}")
    elif ext == ".pptx":
        export_pptx(deck, out_path)
        print(f"Wrote {out_path}")
    else:
        sys.exit(f"error: unsupported output extension '{ext}' (use .html, .pdf, or .pptx)")


if __name__ == "__main__":
    main()
