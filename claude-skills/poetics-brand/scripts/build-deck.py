#!/usr/bin/env python3
"""build-deck.py - render a Poetics slide deck from a deck.json spec.

Runs standalone, no AI required:
    python3 build-deck.py <deck.json> <output.html|output.pdf|output.pptx> [--check] [--sheet]

The output format is chosen by the output file's extension:
  .html  - self-contained deck, zero external dependencies. Opens as a presenter:
           the 1920x1080 slide scales to the window. Keys: arrows / space / page keys
           to move, F fullscreen, N speaker notes, T light/dark, S scroll all slides,
           Home/End. Add ?static to the URL for the stacked, printable view.
  .pdf   - rendered via Playwright/Chromium from the same HTML, one page per slide.
  .pptx  - editable PowerPoint via python-pptx, with speaker notes. A styled
           approximation: the HTML/PDF is the visual truth. Install
           assets/fonts/InstrumentSerif-Regular.ttf and InterTight-Variable.ttf to
           see it as designed.

--check renders every slide with Playwright, writes one PNG per slide to
<output-stem>-check/ beside the output, and reports text that overflows the
slide or crowds the footer. Look at the PNGs before handing a deck over.

--sheet (implies --check) also writes <output-stem>-sheet.png: every slide on one
page, captioned with its page template, for reviewing a deck or choosing templates.

The design rules are in brand-kit/design-system.md (the website's design system
translated to slides). This script applies them; content decisions are the caller's.

deck.json
---------
{
  "kind": "standard",          // standard | curation | proposal
  "mode": "offer",             // proposal only: pitch | offer | tiered | quote
  "title": "Deck title",       // HTML <title>, PPTX metadata
  "client": "Benaki Museum",   // optional, used by templates in the copy
  "number": "2026-014",        // proposal only: shown in the footer as "Nº 2026-014"
  "format": "16:9",            // 16:9 (default) | a4 (portrait; the default for proposals)
  "date": "31 October 2026",   // proposals: shown in the footer as "Created ..." (default: today)
  "disclaimer": "...",         // proposals: footer line (default: any change in scope
                               // requires an adjusted quotation)
  "theme": "light",            // light (default) | dark
  "density": "speaker",        // speaker (default, larger type) | reading (more per slide)
  "footer": "poeticsvr.com",   // footer text, left
  "confidential": true,        // default: "poeticsvr.com · Confidential" in every footer
  "allow_team_names": false,   // true allows the team's own names (only when asked)
  "draft": false,              // true shows [PLACEHOLDERS] highlighted instead of dropping them
  "slides": [ ... ]
}

Every slide takes: "type", "notes" (speaker notes; HTML N key and PPTX notes),
"density" (override the deck's density for one slide),
"overline" (alias "eyebrow"), "note" (small print at the foot of the slide) and
"ground": "paper" | "deep" to override the default ground.

Page templates
--------------
A slide may name a page template instead of a type: {"template": "E3", ...}. The
codes, grouped by what the page does (see PAGE_TEMPLATES below and the skill):
  Open      O1 cover · O2 cover with photo · O3 statement
  Navigate  N1 section divider · N2 agenda
  Explain   E1 claim and rows · E2 claim and terms · E3 claim and columns ·
            E4 claim and photo columns
  Prove     P1 figures · P2 comparison table
  Show      S1 text and photo · S2 photo with claim · S3 full photo · S4 gallery
  Voices    V1 quote · V2 portraits
  Time      T1 steps
  Close     C1 closing (email only)
  Deal      D1 price · D2 acceptance · D3 quote sheet   (proposal decks)
  Program   W1 title card                   (curation decks)
  Explain   E5 text (paragraphs)
  Document  A1 stacked page: several templates on one A4 page
            {"template": "A1", "sections": [{"template": "E5", ...}, {"template": "E2", ...}]}
Filler pages (O3, N1, N2, S3, S4) carry no information of their own: the script lists
them after every render so the user can be asked whether to keep them.

Slide types
-----------
  cover      title, lead (alias subtitle), overline, image, focal
  statement  text, overline, credit, facts: ["Plug-and-play setups", ...]
  section    title, overline
  agenda     title, items: ["...", ...]                          (numbered rows)
  content    title, lead, bullets: ["...", ...]                  (ruled rows)
  terms      title, lead, rows: [{term, text}]                   (label beside text)
  columns    title, lead, columns: [{overline, title, text, image, focal}]
             (alias "cards" with cards: [{title, body}])
  stats      title, overline, stats: [{value, label, sublabel}], note
  split      title, lead, items: [{title, text, note}], image, focal, ratio,
             side: right|left, bleed: true, caption
  timeline   title, steps: [{title, text, when, now: true}]
  quote      quote, name, role (alias attribution), image, focal
  people     title, people: [{name, role, context, image, focal}], shape: lens|pill
  gallery    title, images: [{src, focal}] (or plain paths), caption       (2-8 photos)
  media      title, image, focal, caption, bleed: left|right
  table      title, lead, columns: ["", "Chamber", "Gallery", "Forum"],
             rows: [["Pods", "4", "8", "12"], ...], highlight: 2, highlight_label, note
  price      title, overline, amount: "€ 18,500", basis, includes: [...], validity, note
  work       CURATION DECKS ONLY. title, image, image_credit, focal, pillar, synopsis, fit,
             link: {label, href} (official page or trailer, shown in lens blue),
             facts: [{term, text}]  (duration, year, creator, distributor, languages, age...),
             sources: [{fact, url}]  (not shown; what audit-deck.py checks)
  acceptance title, text, parties: ["For the venue", "For Poetics"]
  closing    title (default "Write to us."). One path only: hello@poeticsvr.com.
  text       title, lead, paragraphs: ["...", ...]
  quote_sheet  (D3) to, title, intro, scope: [{title, text, items: [...]}], capacity, amount,
             basis, validity. The practical one-page proposal.
"lang": "el" writes the renderer's own words (footer, price labels, signatures) in Greek.
  page       sections: [slide, slide, ...]   (A4: stacks templates on one page)

Images are library ids ("photo:<id>", see photo-library.py) or local paths under /tmp, the
brand-kit tree, or the deck.json's own folder.
Text is plain and HTML-escaped. `*asterisks*` from the old emphasis voice are stripped.

Writing rules the script enforces with warnings (see design-system.md): no em dashes,
no visible placeholders (an item containing [CAPS] or [CONFIRM: ...] is dropped unless
"draft": true), no film titles outside curation decks, and the AI-slop phrases listed
in guides/anti-slop.md.
"""

import base64
import html as html_lib
import json
import pathlib
import re
import sys

KIT_ROOT = pathlib.Path(__file__).resolve().parent.parent
FONTS = KIT_ROOT / "assets/fonts"
LOGOS = KIT_ROOT / "assets/logos"

# -- Tokens (design-system.md, "Colour") --------------------------------------------

THEMES = {
    "light": {
        # Decks sit a step deeper than the website's #F4F5F3: off-white, never #FFFFFF.
        "paper": "#EFF0EC", "paper-deep": "#E5E8E3", "shade": "#C5CBC4",
        "ink": "#121614", "ink-soft": "#5B625C", "lens": "#16249C", "lens-deep": "#0E1769",
    },
    "dark": {
        "paper": "#121614", "paper-deep": "#1C211E", "shade": "#3A423D",
        "ink": "#F4F5F3", "ink-soft": "#8A918B", "lens": "#8F9CFF", "lens-deep": "#BCC4FF",
    },
}

SITE = {
    "email": "hello@poeticsvr.com",
}

LATIN = ("U+0000-00FF, U+0100-017F, U+0131, U+0152-0153, U+02BB-02BC, U+02C6, U+02DA, "
         "U+02DC, U+2000-206F, U+2074, U+20AC, U+2122, U+2212, U+2215")
FONT_FACES = [
    ("Instrument Serif", 400, "InstrumentSerif-Regular-latin.woff2", LATIN),
    ("Instrument Serif", 400, "NotoSerifDisplay-w70-greek.woff2",
     "U+0370-0377, U+037A-037F, U+0384-038A, U+038C, U+038E-03A1, U+03A3-03FF"),
    ("Instrument Serif", 400, "NotoSerifDisplay-w70-greek-ext.woff2", "U+1F00-1FFF"),
    ("Inter Tight", 400, "InterTight-400-latin.woff2", LATIN),
    ("Inter Tight", 400, "InterTight-400-greek.woff2", "U+0370-03FF"),
    ("Inter Tight", 500, "InterTight-500-latin.woff2", LATIN),
    ("Inter Tight", 500, "InterTight-500-greek.woff2", "U+0370-03FF"),
]

# The wordmark's O, outer edge, in the wordmark's own coordinates (website mark.ts).
RING_VIEW = "20.068 0 36.791 20"
RING_PATH = ("M 30.102 0 L 46.825 0 C 52.366 0 56.859 4.477 56.859 10 C 56.859 15.523 "
             "52.366 20 46.825 20 L 30.102 20 C 24.56 20 20.068 15.523 20.068 10 "
             "C 20.068 4.477 24.56 0 30.102 0 Z")

# -- Text helpers ---------------------------------------------------------------------

EMPH_RE = re.compile(r"\*([^*]+)\*")
# "[TBC]", "[SQUARE METERS]", "[CONFIRM: venue name]": the website's placeholder convention.
PH_RE = re.compile(r"\[[A-Z][A-Z0-9 _/&+.'’\-]*(?::[^\]]*)?\]")

WARNINGS = []
DRAFT = False


def warn(where, msg):
    WARNINGS.append(f"{where}: {msg}")


def plain(text):
    return EMPH_RE.sub(r"\1", str(text or ""))


def e(text):
    """Escape for HTML. In draft mode placeholders show highlighted."""
    out = html_lib.escape(plain(text))
    if DRAFT:
        out = PH_RE.sub(lambda m: f'<mark class="ph">{m.group(0)}</mark>', out)
    return out


def has_ph(value):
    if isinstance(value, str):
        return bool(PH_RE.search(value))
    if isinstance(value, dict):
        return any(has_ph(v) for v in value.values())
    if isinstance(value, list):
        return any(has_ph(v) for v in value)
    return False


def field(slide, key, where, *aliases):
    """A scalar field, or '' if missing or (outside draft mode) a placeholder."""
    for k in (key, *aliases):
        val = slide.get(k)
        if val:
            if has_ph(val) and not DRAFT:
                warn(where, f'"{k}" is a placeholder and was left out: {val}')
                return ""
            return val
    return ""


def items(slide, key, where, *aliases):
    """A list field with placeholder items dropped (outside draft mode)."""
    for k in (key, *aliases):
        val = slide.get(k)
        if val:
            kept = []
            for it in val:
                if has_ph(it) and not DRAFT:
                    warn(where, f'an item in "{k}" contains a placeholder and was left out: '
                                f'{json.dumps(it, ensure_ascii=False)[:90]}')
                else:
                    kept.append(it)
            return kept
    return []


# -- Copy lint (guides/anti-slop.md) ------------------------------------------------

SLOP_WORDS = [
    "delve", "seamless", "seamlessly", "elevate", "unlock", "unleash", "game-changing",
    "game changer", "revolutionize", "revolutionary", "cutting-edge", "next-level",
    "world-class", "best-in-class", "synergy", "leverage", "empower", "transformative",
    "tapestry", "testament", "pivotal", "vibrant", "nestled", "in today's", "ever-evolving",
    "landscape of", "harness", "robust", "innovative solution", "immersive revolution",
    "take it to the next level", "boasts", "stands as", "serves as",
]
SLOP_RE = re.compile(r"\b(" + "|".join(re.escape(w) for w in SLOP_WORDS) + r")\b", re.I)
NEG_PARALLEL_RE = re.compile(
    r"\b(not (just|only|merely) [^.]{1,60}\bbut\b|it'?s not [^.]{1,60}[,;] it'?s\b|"
    r"more than just\b)", re.I)
EMOJI_RE = re.compile("[\U0001F300-\U0001FAFF☀-➿]")


def lint_text(where, text):
    if not isinstance(text, str) or not text:
        return
    if "—" in text:
        warn(where, f"em dash: use a colon, comma, full stop or parentheses: {text[:70]}")
    m = SLOP_RE.search(text)
    if m:
        warn(where, f'stock phrase "{m.group(0)}": say the concrete thing instead')
    if NEG_PARALLEL_RE.search(text):
        warn(where, f'"not X, but Y" construction: state Y plainly: {text[:70]}')
    if "!" in text:
        warn(where, "exclamation mark: the brand doesn't shout")
    if EMOJI_RE.search(text):
        warn(where, "emoji in slide text")


def lint_heading(where, text):
    words = [w for w in re.findall(r"[A-Za-z][\w'’-]*", PH_RE.sub("", plain(text)))]
    small = {"a", "an", "and", "as", "at", "but", "by", "for", "in", "of", "on", "or",
             "the", "to", "with", "vs"}
    if len(words) >= 4:
        caps = [w for w in words[1:] if w.lower() not in small]
        upper = sum(w[0].isupper() for w in caps)
        # Three or more capitalised words: a client or venue name alone doesn't trip it.
        if caps and upper >= 3 and upper / len(caps) > 0.6:
            warn(where, f"title case in a heading: use sentence case: {text}")


def walk_strings(value):
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for k, v in value.items():
            if k not in ("image", "src", "href", "type", "focal", "side", "shape", "ratio"):
                yield from walk_strings(v)
    elif isinstance(value, list):
        for v in value:
            yield from walk_strings(v)


# -- Assets -------------------------------------------------------------------------

def b64(path: pathlib.Path) -> str:
    return base64.b64encode(path.read_bytes()).decode()


def font_face_css():
    out = []
    for family, weight, filename, urange in FONT_FACES:
        path = FONTS / filename
        if path.exists():
            out.append(
                f"@font-face {{ font-family: '{family}'; font-weight: {weight}; "
                f"font-style: normal; font-display: block; unicode-range: {urange}; "
                f"src: url('data:font/woff2;base64,{b64(path)}') format('woff2'); }}")
    return "\n".join(out)


def wordmark_svg(cls="wordmark"):
    path = LOGOS / "wordmark-mono.svg"
    if not path.exists():
        return ""
    svg = path.read_text(encoding="utf-8")
    svg = re.sub(r'\s(width|height)="[^"]*"', "", svg, count=2)
    svg = svg.replace('fill="#121614"', 'fill="currentColor"')
    return svg.replace("<svg ", f'<svg class="{cls}" role="img" aria-label="Poetics" ', 1)


def ring_svg():
    return (f'<svg class="ring" viewBox="{RING_VIEW}" aria-hidden="true">'
            f'<path d="{RING_PATH}" fill="none" stroke="currentColor" stroke-width="1" '
            f'vector-effect="non-scaling-stroke"/></svg>')


_ALLOWED_EXTRA = []


def _load_brand_assets():
    import importlib.util as _ilu
    spec = _ilu.spec_from_file_location("_brand_assets", pathlib.Path(__file__).parent / "_brand_assets.py")
    mod = _ilu.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _load_photo_library():
    import importlib.util as _ilu
    spec = _ilu.spec_from_file_location("photo_library", pathlib.Path(__file__).parent / "photo-library.py")
    mod = _ilu.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def resolve_image(src: str) -> str:
    """A local image as a data URI, via the shared allowlist (/tmp, the brand-kit tree,
    and the deck.json's folder). '' when it can't be resolved."""
    if not src:
        return ""
    if src.startswith("data:image/"):
        return src
    if src.startswith("photo:"):   # the Poetics photo library (scripts/photo-library.py)
        lib = _load_photo_library()
        path = lib.resolve(src[6:])
        if not path:
            return ""
        ba = _load_brand_assets()
        return ba.resolve_local_image(str(path), (lib.library_root().resolve(),))
    if _ALLOWED_EXTRA and not pathlib.Path(src).expanduser().is_absolute():
        src = str(_ALLOWED_EXTRA[0] / src)   # relative to the deck.json
    ba = _load_brand_assets()
    roots = tuple(ba.make_allowed_roots(__file__)) + tuple(_ALLOWED_EXTRA)
    return ba.resolve_local_image(src, roots)


FOCAL = {
    "center": "50% 50%", "top": "50% 0%", "bottom": "50% 100%", "left": "0% 50%",
    "right": "100% 50%", "top-left": "0% 0%", "top-right": "100% 0%",
    "bottom-left": "0% 100%", "bottom-right": "100% 100%",
}


def aperture(src, where, *, cls="", focal="center", style=""):
    """A photograph in an aperture. Missing -> an empty paper-deep frame, never text."""
    uri = resolve_image(src) if src else ""
    if src and not uri:
        warn(where, f"image not found or outside the allowed folders: {src}")
    if not src:
        warn(where, "no photograph: an empty frame renders in its place")
    pos = FOCAL.get(focal or "center", "50% 50%")
    inner = (f'<img src="{uri}" alt="" style="object-position:{pos}">' if uri else "")
    return f'<div class="ap {cls}" style="{style}">{inner}</div>'


# -- CSS ------------------------------------------------------------------------------

def theme_vars(name):
    return "; ".join(f"--{k}: {v}" for k, v in THEMES[name].items())


CSS = """
* { box-sizing: border-box; margin: 0; padding: 0; }
:root { --display: 'Instrument Serif', 'Iowan Old Style', Georgia, serif;
        --sans: 'Inter Tight', ui-sans-serif, system-ui, -apple-system, 'Helvetica Neue', Arial, sans-serif; }
.deck[data-theme='light'] { __LIGHT__; color-scheme: light; }
.deck[data-theme='dark'] { __DARK__; color-scheme: dark; }
html, body { background: #121614; }
body { font-family: var(--sans); -webkit-font-smoothing: antialiased; }
.deck { background: var(--paper-deep); }

/* The slide: 1920 x 1080. Margins 120; content from 192 to 168 above the foot. */
.slide {
  position: relative; width: 1920px; height: 1080px; overflow: hidden;
  background: var(--paper); color: var(--ink);
  padding: 192px 120px 168px; display: flex; flex-direction: column;
  font-size: 26px; line-height: 1.6; font-kerning: normal;
}
.slide.deep { background: var(--paper-deep); }
.wordmark { position: absolute; top: 72px; left: 120px; height: 30px; width: auto; color: var(--ink); }
.foot { position: absolute; left: 120px; right: 120px; bottom: 56px; display: flex;
        justify-content: space-between; font-size: 21px; line-height: 1.5; color: var(--ink-soft); }
.small-print { position: absolute; left: 120px; right: 120px; bottom: 104px; font-size: 21px;
               line-height: 1.5; color: var(--ink-soft); max-width: 50ch; }
mark.ph { background: var(--paper-deep); color: var(--ink-soft); outline: 1px dashed var(--shade); }

/* Type roles */
.overline { font-size: 21px; line-height: 1.5; color: var(--ink-soft); margin-bottom: 24px; }
.h { font-family: var(--display); font-weight: 400; letter-spacing: -0.02em; text-wrap: balance; }
.h-xl { font-size: 156px; line-height: 1; max-width: 24ch; }
.h-l  { font-size: 102px; line-height: 1.05; max-width: 18ch; }
.h-m  { font-size: 66px; line-height: 1.1; max-width: 36ch; }
.lead { font-size: 32px; line-height: 1.5; max-width: 50ch; margin-top: 48px; text-wrap: pretty; }
.soft { color: var(--ink-soft); }
.t-title { font-size: 36px; line-height: 1.25; font-weight: 500; }
.t-text { font-size: var(--text); line-height: var(--lh); text-wrap: pretty; }
.t-cap { font-size: 21px; line-height: 1.5; color: var(--ink-soft); }
.deck { --text: 32px; --lh: 1.5; }
.deck.reading, .slide.reading { --text: 26px; --lh: 1.6; }
.slide.speaker { --text: 32px; --lh: 1.5; }
.content { margin-top: 96px; flex: 1; min-height: 0; display: flex; flex-direction: column; }
.content.tight { margin-top: 72px; }
.bottom { justify-content: flex-end; }
.center-y { justify-content: center; }

/* Ruled rows: hairlines, 30px row padding, a closing rule. */
.ruled { list-style: none; border-bottom: 1px solid var(--shade); }
.ruled > * { border-top: 1px solid var(--shade); padding: 30px 0; break-inside: avoid; }
.ruled.compact > * { padding: 18px 0; }
.ruled.cols-2 { column-count: 2; column-gap: 36px; border-bottom: 0; }
.ruled.cols-2 > * { display: block; }
.ruled.numbered > * { display: grid; grid-template-columns: 72px minmax(0, 1fr); align-items: baseline; }
.ruled.terms > * { display: grid; grid-template-columns: 420px minmax(0, 1fr); column-gap: 36px; align-items: baseline; }
.ruled .num { font-size: 21px; color: var(--ink-soft); font-variant-numeric: tabular-nums; }
.row-title { font-size: var(--text); line-height: var(--lh); font-weight: 500; }
.row-text { font-size: 26px; line-height: 1.6; color: var(--ink-soft); margin-top: 12px; }
.row-note { font-size: 21px; color: var(--ink-soft); margin-top: 6px; }
.term { font-weight: 500; }

/* Apertures */
.ap { position: relative; overflow: hidden; background: var(--paper-deep); border-radius: 48px; }
.slide.deep .ap { background: var(--paper); }
.ap img { position: absolute; inset: 0; width: 100%; height: 100%; object-fit: cover; display: block; }
.ap.pill { border-radius: 9999px; }
.ap.bleed-right { border-radius: 48px 0 0 48px; margin-right: -120px; }
.ap.bleed-left { border-radius: 0 48px 48px 0; margin-left: -120px; }

/* Cover */
.cover .body { flex: 1; display: flex; flex-direction: column; justify-content: flex-end; }
.cover.with-image { flex-direction: row; gap: 36px; padding-right: 0; }
.cover.with-image .body { flex: 0 0 980px; }
.cover.with-image .cover-media { flex: 1; }   /* starts with the content (192px), never under the wordmark */
.cover .lead { color: var(--ink-soft); }

/* Statement, section */
.statement .facts { display: flex; flex-wrap: wrap; gap: 12px 48px; margin-top: 48px; }
.statement .facts span { font-size: 26px; color: var(--ink-soft); }
.statement .credit { margin-top: 48px; font-size: 26px; color: var(--ink-soft); }

/* Columns, steps */
.grid { display: grid; column-gap: 36px; row-gap: 72px; }
.col { border-top: 1px solid var(--shade); padding-top: 30px; min-width: 0; }
.col.has-image { border-top: 0; padding-top: 0; display: flex; flex-direction: column; }
.col .ap { aspect-ratio: 4 / 3; margin-bottom: 24px; }
.col .overline { margin-bottom: 6px; }
.col .t-title + .row-text { margin-top: 18px; }
.step .when { font-size: 21px; color: var(--ink-soft); margin-top: 18px; }
.step .num { font-size: 21px; color: var(--ink-soft); font-variant-numeric: tabular-nums; margin-bottom: 18px; display: block; }
.step.now .num { color: var(--ink); font-weight: 500; }
.timeline .col { border-top-color: var(--shade); }
.timeline .col.now { border-top: 2px solid var(--ink); padding-top: 29px; }

/* Stats */
.stats .col { padding-top: 30px; }
.stat-value { font-family: var(--display); font-weight: 400; letter-spacing: -0.03em; line-height: 1; }
.stat-value.feature { font-size: 144px; }
.stat-value.compact { font-size: 102px; letter-spacing: -0.02em; }
.stat-label { margin-top: 24px; font-size: 26px; line-height: 1.5; color: var(--ink-soft); max-width: 22ch; }
.stat-sub { margin-top: 6px; font-size: 21px; color: var(--ink-soft); max-width: 30ch; }

/* Split */
.split { display: grid; grid-template-columns: minmax(0, 1fr) minmax(0, 1fr); grid-template-rows: minmax(0, 1fr); column-gap: 36px; flex: 1; min-height: 0; }
.split .text { display: flex; flex-direction: column; min-width: 0; }
.split .text .content { margin-top: 72px; }
.split .media { display: flex; flex-direction: column; min-height: 0; }
.split .media .ap { flex: 1; min-height: 0; }
.split .media .t-cap { margin-top: 18px; }
.split.media-left .media { order: -1; }

/* Quote */
.quote-text { font-family: var(--display); font-size: 66px; line-height: 1.15; letter-spacing: -0.02em; max-width: 30ch; text-wrap: balance; }
.quote-text.long { font-family: var(--sans); font-size: 32px; line-height: 1.5; letter-spacing: 0; max-width: 50ch; }
.speaker { display: flex; align-items: center; gap: 24px; margin-top: 72px; }
.speaker .ap { width: 72px; aspect-ratio: 5 / 9; border-radius: 9999px; flex: none; }
.speaker .name { font-weight: 500; }
.quote-wrap { flex: 1; display: flex; flex-direction: column; justify-content: center; }

/* People */
.people { display: grid; column-gap: 36px; flex: 1; min-height: 0; }
.person { display: flex; flex-direction: column; min-height: 0; }
.person .ap { flex: 1; min-height: 0; aspect-ratio: 3 / 4; align-self: flex-start; max-width: 100%; }
.person .who { margin-top: 24px; min-height: 100px; }  /* name + two lines, so portraits align */
.person .name { font-weight: 500; font-size: 26px; line-height: 1.4; }
.person .role, .person .ctx { font-size: 21px; line-height: 1.5; color: var(--ink-soft); }

/* Gallery, media */
.gallery { display: grid; gap: 36px; flex: 1; min-height: 0; }
.gallery .ap { min-height: 0; }
.media-slide .frame { flex: 1; min-height: 0; display: flex; flex-direction: column; }
.media-slide .frame .ap { flex: 1; min-height: 0; }
.media-slide .t-cap { margin-top: 18px; }

/* Table */
table.ruled-table { width: 100%; border-collapse: separate; border-spacing: 0; font-size: 26px; line-height: 1.5; }
.ruled-table th { text-align: left; font-weight: 400; font-size: 21px; color: var(--ink-soft); padding: 0 24px 18px 0; vertical-align: bottom; }
.ruled-table td { border-top: 1px solid var(--shade); padding: 24px 24px 24px 0; vertical-align: top; }
.ruled-table tr:last-child td { border-bottom: 1px solid var(--shade); }
.ruled-table td:first-child { font-weight: 500; }
.ruled-table .hl { background: var(--paper-deep); padding-left: 24px; }
.slide.deep .ruled-table .hl { background: var(--paper); }
.ruled-table th.hl { border-radius: 18px 18px 0 0; padding-top: 18px; color: var(--ink); }
.ruled-table tr:last-child td.hl { border-radius: 0 0 18px 18px; }

/* Price */
.price { display: grid; grid-template-columns: minmax(0, 7fr) minmax(0, 5fr); column-gap: 96px; flex: 1; align-items: start; }
.price .amount { font-family: var(--display); font-size: 144px; line-height: 1; letter-spacing: -0.03em; }
.price .basis { font-size: 32px; color: var(--ink-soft); margin-top: 24px; }
.price .validity { font-size: 26px; color: var(--ink-soft); margin-top: 48px; max-width: 40ch; }
.price .incl-label { font-size: 21px; color: var(--ink-soft); margin-bottom: 18px; }
.price .ruled > * { padding: 18px 0; font-size: 26px; }

/* Work (curation) */
.work { display: grid; grid-template-columns: minmax(0, 7fr) minmax(0, 5fr); column-gap: 72px; flex: 1; min-height: 0; }
.work .still { display: flex; flex-direction: column; min-height: 0; }
.work .still .ap { flex: 1; min-height: 0; }
.work .still .t-cap { margin-top: 18px; }
.work-link { margin-top: 24px; font-size: 26px; }
.work-link a { color: var(--lens); text-decoration: underline; text-decoration-thickness: 1px; text-underline-offset: 6px; }
.work .info { display: flex; flex-direction: column; min-height: 0; }
.work .synopsis { margin-top: 36px; font-size: 26px; line-height: 1.6; }
.work .fit { margin-top: 24px; font-size: 26px; line-height: 1.6; color: var(--ink-soft); }
.work .facts { margin-top: auto; font-size: 21px; }
.work .facts > * { padding: 12px 0; grid-template-columns: 180px minmax(0, 1fr); }
.work .facts .term { font-weight: 400; color: var(--ink-soft); }


/* Acceptance */
.parties { display: grid; grid-template-columns: 1fr 1fr; column-gap: 96px; margin-top: auto; }
.party .label { font-weight: 500; margin-bottom: 72px; }
.sig { border-top: 1px solid var(--ink-soft); padding-top: 12px; font-size: 21px; color: var(--ink-soft); margin-bottom: 72px; }

/* Closing: the band (a wide pill on paper-deep) with the Ring crossing it. */
/* Closing: one path, the email. The band (a wide pill on paper-deep) with the Ring crossing it. */
.closing .band { position: relative; isolation: isolate; overflow: hidden; background: var(--paper-deep);
  border-radius: 9999px; min-height: 480px; padding: 96px 192px; display: flex; flex-direction: column;
  justify-content: center; align-items: flex-start; gap: 36px; margin-top: auto; margin-bottom: auto; }
.closing .band .h { max-width: 18ch; }
.closing .email { font-family: var(--display); font-size: 66px; line-height: 1.1; letter-spacing: -0.02em;
  color: var(--lens); text-decoration: underline; text-decoration-thickness: 2px; text-underline-offset: 12px; }
.closing .ring { position: absolute; z-index: -1; height: 768px; right: -240px; bottom: -240px; color: var(--shade); }

/* Text (E5) and stacked pages (A1) */
.prose { max-width: 50ch; font-size: var(--text); line-height: 1.6; }
.prose p + p { margin-top: 24px; }
.flow { display: flex; flex-direction: column; flex: 1; min-height: 0; }
.flow > .sec + .sec { margin-top: 120px; }
.flow .sec { display: flex; flex-direction: column; }
.flow .sec .content { margin-top: 48px; flex: none; }
.flow .sec .split { flex: none; }

/* Quote sheet (D3): the practical proposal page. */
.qs-intro { margin-top: 48px; max-width: 62ch; }
.qs-scope .subs { list-style: none; margin-top: 6px; }
.qs-scope .subs li { font-size: 26px; line-height: 1.6; color: var(--ink-soft); padding-left: 24px; }
.qs-capacity { margin-top: 48px; font-size: var(--text); line-height: 1.6; }
.qs-price { margin-top: 72px; }
.qs-price .amount { font-family: var(--display); font-size: 102px; line-height: 1; letter-spacing: -0.02em; }
.qs-price .basis { font-size: 26px; color: var(--ink-soft); margin-top: 18px; }
.qs-price .validity { font-size: 21px; color: var(--ink-soft); margin-top: 12px; }
.qs-price .incl-label { font-size: 21px; color: var(--ink-soft); margin-bottom: 12px; }

/* Footer: a proposal adds its scope disclaimer as a first line. */
.foot { flex-wrap: wrap; row-gap: 6px; }
.foot .disclaimer { flex-basis: 100%; }
.slide.has-disclaimer { padding-bottom: 216px; }
.slide.has-disclaimer .small-print { bottom: 152px; }

/* A4 portrait: 1480 x 2092, the same type tokens on a reading page (body lands near 10.5pt). */
.deck.a4 .slide { width: 1480px; height: 2092px; padding: 192px 120px 216px; }
.deck.a4 .slide.has-disclaimer { padding-bottom: 264px; }
.deck.a4 .cover.with-image { flex-direction: column; padding-right: 120px; gap: 0; }
.deck.a4 .cover.with-image .cover-media { order: -1; flex: 0 0 900px; margin: 0 -120px 96px 0; }
.deck.a4 .cover.with-image .body { flex: 1 1 auto; }
.deck.a4 .split { grid-template-columns: minmax(0, 1fr); grid-template-rows: auto minmax(0, 1fr); row-gap: 72px; }
.deck.a4 .split.media-left .media { order: 0; }
.deck.a4 .price { grid-template-columns: minmax(0, 1fr) minmax(0, 1fr); column-gap: 72px; }
.deck.a4 .timeline .grid { grid-template-columns: minmax(0, 1fr) !important; row-gap: 0; }
.deck.a4 .timeline .col { display: grid; grid-template-columns: 96px minmax(0, 1fr) 260px;
  grid-template-areas: "num title when" "num text when"; column-gap: 36px; padding: 30px 0; align-items: baseline; }
.deck.a4 .timeline .grid > .col:last-child { border-bottom: 1px solid var(--shade); }
.deck.a4 .timeline .col.now { padding-top: 29px; }
.deck.a4 .step .num { grid-area: num; margin: 0; }
.deck.a4 .step .t-title { grid-area: title; font-size: 26px; }
.deck.a4 .step .row-text { grid-area: text; margin-top: 6px; }
.deck.a4 .step .when { grid-area: when; margin: 0; text-align: right; }
.deck.a4 .closing .band { border-radius: 48px; min-height: 0; padding: 120px 96px; }
.deck.a4 .closing .ring { height: 640px; right: -200px; bottom: -220px; }
.deck.a4 .prose { max-width: 62ch; }
html:not(.present) .deck.a4 .stage { width: 1480px; margin: 0 auto; padding: 24px 0; }

/* Presenter. JavaScript adds .present unless the URL has ?static or the page is printed. */
html.present, html.present body { height: 100%; overflow: hidden; }
html.present body { background: var(--letterbox, #1C211E); }
html.present .deck { position: fixed; inset: 0; background: transparent; }
html.present .stage { position: absolute; left: 0; top: 0; width: 1920px; height: 1080px; transform-origin: 0 0; }
html.present .slide { position: absolute; inset: 0; visibility: hidden; }
html.present .slide.current { visibility: visible; }
.notes { display: none; }
.notes-panel { display: none; position: fixed; left: 0; right: 0; bottom: 0; max-height: 40vh; overflow: auto;
  padding: 16px 24px; font: 17px/1.6 var(--sans); background: var(--paper-deep); color: var(--ink);
  border-top: 1px solid var(--shade); z-index: 10; }
html.show-notes .notes-panel { display: block; }
.notes-panel .count { color: var(--ink-soft); font-size: 14px; }
html:not(.present) .slide + .slide { margin-top: 24px; }

@media print {
  html, body { background: none; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
  html.present, html.present body { height: auto; overflow: visible; }
  html.present .deck, html.present .stage { position: static; transform: none !important; }
  html.present .slide { position: relative; visibility: visible; }
  .slide + .slide { margin-top: 0 !important; }
  html:not(.present) .deck.a4 .stage { padding: 0 !important; zoom: 1 !important; }
  .slide { break-after: page; }
  .notes-panel { display: none !important; }
}
"""

JS = r"""
(function () {
  var html = document.documentElement;
  var params = new URLSearchParams(location.search);
  if (params.has('static')) return;
  var deck = document.querySelector('.deck');
  if (deck.classList.contains('a4')) {
    var st = document.querySelector('.stage');
    var zoom = function () { st.style.zoom = Math.min(1, (innerWidth - 48) / 1480); };
    addEventListener('resize', zoom); zoom();
    return;
  }
  var stage = document.querySelector('.stage');
  var slides = Array.prototype.slice.call(document.querySelectorAll('.slide'));
  var panel = document.querySelector('.notes-panel');
  var i = 0;
  html.classList.add('present');
  function fit() {
    var s = Math.min(innerWidth / 1920, innerHeight / 1080);
    var x = (innerWidth - 1920 * s) / 2, y = (innerHeight - 1080 * s) / 2;
    stage.style.transform = 'translate(' + x + 'px,' + y + 'px) scale(' + s + ')';
  }
  function show(n) {
    i = Math.max(0, Math.min(slides.length - 1, n));
    slides.forEach(function (el, k) { el.classList.toggle('current', k === i); });
    var notes = slides[i].querySelector('.notes');
    panel.innerHTML = '<div class="count">' + (i + 1) + ' of ' + slides.length + '</div>' +
      (notes ? notes.innerHTML : '<p>No notes for this slide.</p>');
    if (location.hash !== '#' + (i + 1)) history.replaceState(null, '', '#' + (i + 1));
  }
  function letterbox() {
    html.style.setProperty('--letterbox', getComputedStyle(deck).getPropertyValue('--paper-deep'));
  }
  addEventListener('resize', fit);
  addEventListener('keydown', function (ev) {
    if (ev.metaKey || ev.ctrlKey || ev.altKey) return;
    var k = ev.key;
    if (k === 'ArrowRight' || k === 'ArrowDown' || k === 'PageDown' || k === ' ' || k === 'Enter') show(i + 1);
    else if (k === 'ArrowLeft' || k === 'ArrowUp' || k === 'PageUp' || k === 'Backspace') show(i - 1);
    else if (k === 'Home') show(0);
    else if (k === 'End') show(slides.length - 1);
    else if (k === 'f' || k === 'F') { if (document.fullscreenElement) document.exitFullscreen(); else html.requestFullscreen(); }
    else if (k === 'n' || k === 'N') html.classList.toggle('show-notes');
    else if (k === 't' || k === 'T') { deck.dataset.theme = deck.dataset.theme === 'dark' ? 'light' : 'dark'; letterbox(); }
    else if (k === 's' || k === 'S') { location.search = 'static'; }
    else return;
    ev.preventDefault();
  });
  stage.addEventListener('click', function (ev) {
    if (ev.target.closest('a')) return;
    show(ev.clientX < innerWidth / 3 ? i - 1 : i + 1);
  });
  var tx = null;
  addEventListener('touchstart', function (ev) { tx = ev.touches[0].clientX; }, { passive: true });
  addEventListener('touchend', function (ev) {
    if (tx === null) return;
    var dx = ev.changedTouches[0].clientX - tx; tx = null;
    if (Math.abs(dx) > 40) show(dx < 0 ? i + 1 : i - 1);
  });
  letterbox(); fit();
  show((parseInt(location.hash.slice(1), 10) || 1) - 1);
})();
"""


# -- Page templates (the catalog in the deck skill) ---------------------------------

PAGE_TEMPLATES = {
    "O1": ("Open", "Cover", "cover"),
    "O2": ("Open", "Cover with photo", "cover"),
    "O3": ("Open", "Statement", "statement"),
    "N1": ("Navigate", "Section divider", "section"),
    "N2": ("Navigate", "Agenda", "agenda"),
    "E1": ("Explain", "Claim and rows", "content"),
    "E2": ("Explain", "Claim and terms", "terms"),
    "E3": ("Explain", "Claim and columns", "columns"),
    "E4": ("Explain", "Claim and photo columns", "columns"),
    "P1": ("Prove", "Figures", "stats"),
    "P2": ("Prove", "Comparison table", "table"),
    "S1": ("Show", "Text and photo", "split"),
    "S2": ("Show", "Photo with claim", "media"),
    "S3": ("Show", "Full photo", "media"),
    "S4": ("Show", "Gallery", "gallery"),
    "V1": ("Voices", "Quote", "quote"),
    "V2": ("Voices", "Portraits", "people"),
    "T1": ("Time", "Steps", "timeline"),
    "C1": ("Close", "Closing", "closing"),
    "D1": ("Deal", "Price", "price"),
    "D2": ("Deal", "Acceptance", "acceptance"),
    "W1": ("Program", "Title card", "work"),
    "E5": ("Explain", "Text", "text"),
    "A1": ("Document", "Stacked page", "page"),
    "D3": ("Deal", "Quote sheet", "quote_sheet"),
}

# Pages that carry no information of their own. Kept only when the user says so.
FILLER = {"O3", "N1", "N2", "S3", "S4"}
FILLER_TYPES = {"statement": "O3", "section": "N1", "agenda": "N2", "gallery": "S4"}


def filler_code(s):
    if s.get("template") in FILLER:
        return s["template"]
    if s.get("template"):
        return None
    if s.get("type") == "media" and not s.get("title"):
        return "S3"
    return FILLER_TYPES.get(s.get("type"))


def resolve_template(s, where):
    """Fill in the slide type from a page-template code, and check the code's needs."""
    code = s.get("template")
    if not code:
        return s
    if code not in PAGE_TEMPLATES:
        warn(where, f'unknown page template "{code}"')
        return s
    s = dict(s)
    want = PAGE_TEMPLATES[code][2]
    if s.get("type") and s["type"] != want and not (want == "columns" and s["type"] == "cards"):
        warn(where, f'template {code} is a "{want}" page but the slide says type "{s["type"]}"')
    s["type"] = s.get("type") or want
    if code == "O2" and not s.get("image"):
        warn(where, "template O2 needs an image")
    if code == "E4" and not all(c.get("image") for c in s.get("columns", [])):
        warn(where, "template E4 needs a photo on every column")
    if code == "S3" and s.get("title"):
        warn(where, "template S3 is a photo with no heading: use S2 for a photo with a claim")
    return s


def template_label(s):
    code = s.get("template")
    if code in PAGE_TEMPLATES:
        return f"{code} · {PAGE_TEMPLATES[code][1]}"
    return s.get("type", "content")


# -- Deck rules (guides/anti-slop.md, "Poetics deck rules") -------------------------

HOUSE_EMAIL = "hello@poeticsvr.com"
TEAM_NAMES = ["Thanos", "Byron", "Ilias", "Dimitriou"]
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
TAGLINE_RE = re.compile(
    r"\b(where [\w\s]{2,30}? meets?\b|redefin\w*|reimagin\w*|the future of\b|step into\b|"
    r"like never before|beyond imagination|experience the magic|magical?\b|unforgettable|"
    r"breathtaking|dive into|immerse yourself|a whole new (world|way)|the art of\b|"
    r"one of a kind|second to none|limitless|endless possibilities|"
    # chatty phrases that read as AI copy
    r"ready when you are|as we understood it|at a glance|shall we|let'?s\b|let us\b|here'?s\b|"
    r"imagine\b|picture this|what if\b|the best part|simply put|a closer look|dive\b|"
    r"we'?d love|we would love|don'?t hesitate|look no further|solutions?\b|transform your space|"
    r"journey|all in one place|take a look|rest assured|peace of mind|hassle-free|worry-free|"
    r"sit back|tailor-made|bespoke|next-generation|game.?changer)", re.I)
TEAM_RE = re.compile(r"\b(" + "|".join(TEAM_NAMES) + r")\b")

# Word limits per field: past these, the slide is saying too much.
WORD_LIMITS = {
    "title": 12, "text": 18, "lead": 22, "intro": 50, "capacity": 20, "subtitle": 22, "credit": 14, "quote": 45,
    "synopsis": 45, "fit": 25, "caption": 16, "note": 18, "validity": 18, "contact": 12,
}
ITEM_LIMITS = {"bullets": 14, "items": 14, "columns": 18, "cards": 18, "steps": 18, "rows": 16, "scope": 20,
               "includes": 8, "paragraphs": 60}


def words(text):
    return len(re.findall(r"\S+", PH_RE.sub("x", plain(text))))


def lint_heading_style(where, s):
    """Website style: a heading is a plain sentence ending in a full stop, never a question."""
    if s.get("type") in ("work",):
        return
    for key in ("title",) + (("text",) if s.get("type") == "statement" else ()):
        val = s.get(key)
        if not isinstance(val, str) or not val.strip() or PH_RE.fullmatch(val.strip()):
            continue
        v = plain(val).strip()
        if v.endswith("?"):
            warn(where, f"question as a heading: state it instead: {v}")
        elif v[-1] not in ".!:" and not v.endswith("]"):
            warn(where, f'heading without a full stop (website style: "We run it ourselves."): {v}')


def lint_rules(where, s, deck, density):
    lint_heading_style(where, s)
    for text in walk_strings({k: v for k, v in s.items() if k != "notes"}):
        for m in EMAIL_RE.findall(text):
            if m.lower() != HOUSE_EMAIL:
                warn(where, f"email {m}: decks use {HOUSE_EMAIL} only")
        m = TAGLINE_RE.search(text)
        if m:
            warn(where, f'sounds like AI copy ("{m.group(0)}"): say it plainly, the way the website does')
        if not deck.get("allow_team_names"):
            m = TEAM_RE.search(text)
            if m:
                warn(where, f'names the team ("{m.group(0)}"): leave the makers out unless asked')
    for key, limit in WORD_LIMITS.items():
        val = s.get(key)
        if isinstance(val, str) and words(val) > limit:
            warn(where, f'"{key}" is {words(val)} words (keep it under {limit}): cut it down')
    bonus = 6 if density == "reading" else 0
    for key, limit in ITEM_LIMITS.items():
        for it in s.get(key) or []:
            if isinstance(it, list):   # a table row: each cell is its own short line
                continue
            t = it if isinstance(it, str) else (it.get("text") or it.get("body") or "")
            if isinstance(t, str) and words(t) > limit + bonus:
                warn(where, f'an item in "{key}" is {words(t)} words: one short line each: '
                            f'{plain(t)[:50]}...')


LIST_KEYS = ("bullets", "items", "columns", "cards", "stats", "steps", "rows", "people", "includes")


def lint_rhythm(deck):
    slides = deck.get("slides", [])
    with_overline = [k + 1 for k, s in enumerate(slides)
                     if (s.get("overline") or s.get("eyebrow")) and s.get("type") not in ("cover", "title")
                     and s.get("template") not in ("O1", "O2")]
    allowed = max(2, len(slides) // 4)
    if len(with_overline) > allowed:
        warn("deck", f"overlines on {len(with_overline)} slides ({', '.join(map(str, with_overline))}): "
                     f"keep them to {allowed} or fewer, only where they add a fact")
    counts = [len(s.get(key)) for s in slides for key in LIST_KEYS
              if isinstance(s.get(key), list) and len(s.get(key)) > 1]
    threes = counts.count(3)
    if threes >= 3 and threes > len(counts) / 2:
        warn("deck", f"{threes} of {len(counts)} lists have exactly three items: use the true count, "
                     f"not everything comes in threes")


# -- Slide renderers ----------------------------------------------------------------

def overline(slide, where):
    text = field(slide, "overline", where, "eyebrow")
    return f'<div class="overline">{e(text)}</div>' if text else ""


def heading(slide, where, size="m", key="title"):
    text = field(slide, key, where)
    if not text:
        return ""
    lint_heading(where, text)
    n = len(plain(text))
    if size == "m" and n > 62:
        warn(where, f"heading is {n} characters; at display-m that runs past two lines. "
                    f"Shorten it to the one claim the slide makes.")
    return f'<h1 class="h h-{size}">{e(text)}</h1>'


def step_down(text, sizes, limits):
    """Pick the largest display token whose character limit the text fits."""
    n = len(plain(text))
    for size, limit in zip(sizes, limits):
        if n <= limit:
            return size
    return sizes[-1]


def lead(slide, where, cls="lead"):
    text = field(slide, "lead", where, "subtitle")
    return f'<p class="{cls}">{e(text)}</p>' if text else ""


def s_cover(s, w):
    title = field(s, "title", w)
    lint_heading(w, title)
    img = s.get("image")
    size = step_down(title, ["xl", "l"], [40 if not img else 18, 999]) if title else "xl"
    body = (f'<div class="body">{overline(s, w)}<h1 class="h h-{size}">{e(title)}</h1>'
            f'{lead(s, w)}</div>')
    if img:
        media = aperture(img, w, cls="bleed-right cover-media", focal=s.get("focal"))
        return "cover with-image", body + media
    return "cover", body


def s_statement(s, w):
    text = field(s, "text", w, "title")
    size = step_down(text, ["l", "m"], [90, 999])
    facts = items(s, "facts", w)
    facts_html = ('<div class="facts">' + "".join(f"<span>{e(f)}</span>" for f in facts) + "</div>"
                  if facts else "")
    credit = field(s, "credit", w)
    credit_html = f'<p class="credit">{e(credit)}</p>' if credit else ""
    return "statement deep center-y", (f'{overline(s, w)}<p class="h h-{size}">{e(text)}</p>'
                                       f'{facts_html}{credit_html}')


def s_section(s, w):
    title = field(s, "title", w)
    lint_heading(w, title)
    size = step_down(title, ["l", "m"], [40, 999])
    return "section deep bottom", f'{overline(s, w)}<h1 class="h h-{size}">{e(title)}</h1>'


def s_agenda(s, w):
    its = items(s, "items", w)
    cols = " cols-2" if len(its) > 5 else ""
    rows = "".join(f'<li><span class="num">{k + 1:02d}</span><span class="t-text">{e(t)}</span></li>'
                   for k, t in enumerate(its))
    return "agenda", (f'{overline(s, w)}{heading(s, w)}'
                      f'<div class="content"><ol class="ruled numbered{cols}">{rows}</ol></div>')


def s_content(s, w, density):
    its = items(s, "bullets", w)
    if density == "speaker" and len(its) > 5:
        warn(w, f"{len(its)} bullets in a speaker-led deck: split the slide (one idea per slide)")
    for b in its:
        if len(plain(b)) > 140:
            warn(w, f"bullet over 140 characters: move the detail to the notes: {plain(b)[:60]}...")
    cols = " cols-2" if len(its) > 5 else ""
    rows = "".join(f'<li class="t-text">{e(b)}</li>' for b in its)
    return "content-slide", (f'{overline(s, w)}{heading(s, w)}{lead(s, w)}'
                             f'<div class="content"><ul class="ruled{cols}">{rows}</ul></div>')


def s_terms(s, w):
    rows = items(s, "rows", w)
    html = "".join(f'<li><span class="t-text term">{e(r.get("term", ""))}</span>'
                   f'<span class="t-text soft">{e(r.get("text", ""))}</span></li>' for r in rows)
    return "terms-slide", (f'{overline(s, w)}{heading(s, w)}{lead(s, w)}'
                           f'<div class="content"><ul class="ruled terms{" compact" if len(rows) > 5 else ""}">{html}</ul></div>')


FORMAT = "16:9"
LANG = "en"

# Words the renderer writes itself, per deck language ("lang": "en" | "el").
LABELS = {
    "en": {"confidential": "Confidential", "created": "Created", "of": "of", "for": "For",
           "basis": "All-inclusive, excl. VAT", "covers": "What it covers", "price": "Price",
           "sig": ["Name and role", "Signature", "Date"], "parties": ["For the client", "For Poetics"],
           "closing": "Write to us.", "disclaimer": "Any change in scope requires an adjusted quotation."},
    "el": {"confidential": "Εμπιστευτικό", "created": "Ημερομηνία", "of": "από", "for": "Προς:",
           "basis": "Χωρίς ΦΠΑ", "covers": "Περιλαμβάνει", "price": "Τελικό ποσό",
           "sig": ["Ονοματεπώνυμο και ιδιότητα", "Υπογραφή", "Ημερομηνία"],
           "parties": ["Για τον πελάτη", "Για την Poetics"], "closing": "Γράψτε μας.",
           "disclaimer": "Κάθε αλλαγή στο scope απαιτεί αναπροσαρμογή της προσφοράς."},
}


def L(key):
    return LABELS.get(LANG, LABELS["en"])[key]


def grid_cols(n):
    if FORMAT == "a4":
        return {1: 1, 2: 2, 3: 3, 4: 2, 5: 3, 6: 3}.get(n, 2)
    return {1: 1, 2: 2, 3: 3, 4: 4, 5: 3, 6: 3}.get(n, 4)


def s_columns(s, w):
    cols = items(s, "columns", w, "cards")
    if len(cols) > 6:
        warn(w, f"{len(cols)} columns: keep it to six or fewer")
        cols = cols[:6]
    if len(cols) in (3, 4) and not any(c.get("image") for c in cols):
        pass
    n = grid_cols(len(cols))
    out = []
    for c in cols:
        img = c.get("image")
        media = aperture(img, w, focal=c.get("focal")) if img else ""
        ov = f'<div class="overline">{e(c.get("overline"))}</div>' if c.get("overline") else ""
        text = c.get("text") or c.get("body") or ""
        out.append(f'<div class="col{" has-image" if img else ""}">{media}{ov}'
                   f'<div class="t-title">{e(c.get("title", ""))}</div>'
                   + (f'<p class="row-text">{e(text)}</p>' if text else "") + "</div>")
    return "columns-slide", (f'{overline(s, w)}{heading(s, w)}{lead(s, w)}'
                             f'<div class="content"><div class="grid" style="grid-template-columns:'
                             f'repeat({n}, minmax(0, 1fr))">{"".join(out)}</div></div>')


def s_stats(s, w):
    stats = items(s, "stats", w)
    n = len(stats)
    if n > 6:
        warn(w, f"{n} figures: keep the three that support the heading's claim")
        stats = stats[:6]
        n = 6
    size = "feature" if n <= 3 else "compact"
    cols = (n if n <= 4 else 3) if FORMAT != "a4" else (n if n <= 3 else 2)
    out = "".join(
        f'<div class="col"><div class="stat-value {size}">{e(st.get("value", ""))}</div>'
        f'<div class="stat-label">{e(st.get("label", ""))}</div>'
        + (f'<div class="stat-sub">{e(st.get("sublabel"))}</div>' if st.get("sublabel") else "")
        + "</div>" for st in stats)
    return "stats", (f'{overline(s, w)}{heading(s, w)}'
                     f'<div class="content tight"><div class="grid" style="grid-template-columns:'
                     f'repeat({cols}, minmax(0, 1fr)); row-gap: 48px">{out}</div></div>')


def s_split(s, w):
    its = items(s, "items", w)
    rows = "".join(
        f'<li><div class="row-title">{e(it.get("title", ""))}</div>'
        + (f'<p class="row-text">{e(it.get("text") or it.get("body"))}</p>' if (it.get("text") or it.get("body")) else "")
        + (f'<p class="row-note">{e(it.get("note"))}</p>' if it.get("note") else "")
        + "</li>" for it in its)
    side = "left" if s.get("side") == "left" else "right"
    bleed = f" bleed-{side}" if s.get("bleed") else ""
    ratio = s.get("ratio")
    style = f"flex:none; aspect-ratio:{ratio.replace(':', ' / ')}" if ratio else ""
    cap = field(s, "caption", w)
    media = (f'<div class="media">{aperture(s.get("image"), w, cls=bleed.strip(), focal=s.get("focal"), style=style)}'
             + (f'<p class="t-cap">{e(cap)}</p>' if cap else "") + "</div>")
    text = (f'<div class="text">{overline(s, w)}{heading(s, w)}{lead(s, w)}'
            + (f'<div class="content"><ul class="ruled">{rows}</ul></div>' if rows else "") + "</div>")
    return "split-slide", f'<div class="split media-{side}">{text}{media}</div>'


def s_timeline(s, w):
    steps = items(s, "steps", w)
    n = max(len(steps), 1)
    out = "".join(
        f'<div class="col step{" now" if st.get("now") else ""}">'
        f'<span class="num">{"Now" if st.get("now") else f"{k + 1:02d}"}</span>'
        f'<div class="t-title">{e(st.get("title", ""))}</div>'
        + (f'<p class="row-text">{e(st.get("text") or st.get("body"))}</p>' if (st.get("text") or st.get("body")) else "")
        + (f'<p class="when">{e(st.get("when"))}</p>' if st.get("when") else "")
        + "</div>" for k, st in enumerate(steps))
    return "timeline", (f'{overline(s, w)}{heading(s, w)}{lead(s, w)}'
                        f'<div class="content"><div class="grid" style="grid-template-columns:'
                        f'repeat({n}, minmax(0, 1fr))">{out}</div></div>')


def s_quote(s, w):
    q = field(s, "quote", w)
    long = " long" if len(plain(q)) > 160 else ""
    name = field(s, "name", w, "attribution")
    role = field(s, "role", w)
    portrait = aperture(s["image"], w, focal=s.get("focal") or "top") if s.get("image") else ""
    who = (f'<div class="speaker">{portrait}<div><div class="name">{e(name)}</div>'
           + (f'<div class="soft">{e(role)}</div>' if role else "") + "</div></div>") if name else ""
    return "quote deep", (f'<div class="quote-wrap">{overline(s, w)}'
                          f'<blockquote class="quote-text{long}">“{e(q)}”</blockquote>{who}</div>')


def s_people(s, w):
    ppl = items(s, "people", w)
    if len(ppl) > 5:
        warn(w, "more than five people: split across two slides")
        ppl = ppl[:5]
    shape = "pill" if s.get("shape") == "pill" else ""
    out = "".join(
        f'<div class="person">{aperture(p.get("image"), w, cls=shape, focal=p.get("focal") or "center")}'
        f'<div class="who"><div class="name">{e(p.get("name", ""))}</div>'
        + (f'<div class="role">{e(p.get("role"))}</div>' if p.get("role") else "")
        + (f'<div class="ctx">{e(p.get("context"))}</div>' if p.get("context") else "")
        + "</div></div>" for p in ppl)
    head = f'{overline(s, w)}{heading(s, w)}'
    cls = "content tight" if head else "content"
    style = "" if head else ' style="margin-top:0"'
    return "people-slide", (f'{head}<div class="{cls}"{style}><div class="people" style="grid-template-columns:'
                            f'repeat({max(len(ppl), 1)}, minmax(0, 1fr))">{out}</div></div>')


def s_gallery(s, w):
    imgs = s.get("images") or []
    if len(imgs) > 8:
        warn(w, f"{len(imgs)} photographs: the gallery shows eight; the rest were left out")
        imgs = imgs[:8]
    if len(imgs) < 2:
        warn(w, "a gallery needs at least two photographs: use a media slide for one")
    n = len(imgs)
    cols, rows = {2: (2, 1), 3: (3, 1), 4: (2, 2), 5: (3, 2), 6: (3, 2), 7: (4, 2), 8: (4, 2)}.get(n, (1, 1))
    tiles = []
    for k, im in enumerate(imgs):
        src, focal = (im, "center") if isinstance(im, str) else (im.get("src"), im.get("focal"))
        span = "grid-row: span 2;" if n in (5, 7) and k == 0 else ""
        tiles.append(aperture(src, w, focal=focal, style=span))
    head = f'{overline(s, w)}{heading(s, w)}'
    cap = field(s, "caption", w)
    cls = "content tight" if head else "content"
    style = "" if head else ' style="margin-top:0"'
    return "gallery-slide", (f'{head}<div class="{cls}"{style}><div class="gallery" style="grid-template-columns:'
                             f'repeat({cols}, minmax(0, 1fr)); grid-template-rows: repeat({rows}, minmax(0, 1fr))">'
                             f'{"".join(tiles)}</div>' + (f'<p class="t-cap" style="margin-top:18px">{e(cap)}</p>' if cap else "")
                             + "</div>")


def s_media(s, w):
    bleed = s.get("bleed")
    cls = f"bleed-{bleed}" if bleed in ("left", "right") else ""
    cap = field(s, "caption", w)
    head = f'{overline(s, w)}{heading(s, w)}'
    frame = (f'<div class="frame">{aperture(s.get("image"), w, cls=cls, focal=s.get("focal"))}'
             + (f'<p class="t-cap">{e(cap)}</p>' if cap else "") + "</div>")
    cls2 = "content tight" if head else "content"
    style = "" if head else ' style="margin-top:0"'
    return "media-slide", f'{head}<div class="{cls2}"{style}>{frame}</div>'


def s_table(s, w):
    cols = s.get("columns") or []
    rows = items(s, "rows", w)
    hl = s.get("highlight")
    def cell(tag, k, v):
        c = ' class="hl"' if hl is not None and k == hl else ""
        return f"<{tag}{c}>{e(v)}</{tag}>"
    head = ""
    if cols:
        labels = list(cols)
        if hl is not None and s.get("highlight_label") and hl < len(labels):
            labels[hl] = f'{plain(labels[hl])} · {plain(s["highlight_label"])}'
        head = "<thead><tr>" + "".join(cell("th", k, v) for k, v in enumerate(labels)) + "</tr></thead>"
    body = "<tbody>" + "".join("<tr>" + "".join(cell("td", k, v) for k, v in enumerate(r)) + "</tr>"
                               for r in rows) + "</tbody>"
    return "table-slide", (f'{overline(s, w)}{heading(s, w)}{lead(s, w)}'
                           f'<div class="content tight"><table class="ruled-table">{head}{body}</table></div>')


def s_price(s, w):
    amount = field(s, "amount", w)
    if not amount:
        warn(w, "price slide without an amount: the slide isn't ready (never guess a price)")
    basis = field(s, "basis", w) or L("basis")
    incl = items(s, "includes", w)
    validity = field(s, "validity", w)
    incl_html = (f'<div><p class="incl-label">{e(L("covers"))}</p><ul class="ruled">'
                 + "".join(f"<li>{e(x)}</li>" for x in incl) + "</ul></div>") if incl else ""
    return "price-slide", (f'{overline(s, w)}{heading(s, w)}<div class="content"><div class="price">'
                           f'<div><div class="amount">{e(amount)}</div><p class="basis">{e(basis)}</p>'
                           + (f'<p class="validity">{e(validity)}</p>' if validity else "")
                           + f"</div>{incl_html}</div></div>")


def s_quote_sheet(s, w):
    """The practical proposal page: who it's for, what's in scope, what it costs. One A4 page."""
    if FORMAT != "a4":
        warn(w, 'the quote sheet (D3) is an A4 page: set "format": "a4" (the default for proposals)')
    to = field(s, "to", w)
    intro = field(s, "intro", w, "lead")
    scope = items(s, "scope", w)
    rows = []
    for it in scope:
        subs = [x for x in (it.get("items") or []) if DRAFT or not has_ph(x)]
        rows.append(f'<li><div class="row-title">{e(it.get("title", ""))}</div>'
                    + (f'<p class="row-text">{e(it.get("text"))}</p>' if it.get("text") else "")
                    + ('<ul class="subs">' + "".join(f"<li>{e(x)}</li>" for x in subs) + "</ul>" if subs else "")
                    + "</li>")
    capacity = field(s, "capacity", w)
    amount = field(s, "amount", w)
    if not amount:
        warn(w, "quote sheet without an amount: the page isn't ready (never guess a price)")
    basis = field(s, "basis", w) or L("basis")
    validity = field(s, "validity", w)
    to_html = f'<div class="overline">{e(L("for"))} {e(to)}</div>' if to else overline(s, w)
    return "quote-sheet", (
        f'{to_html}{heading(s, w)}'
        + (f'<div class="prose qs-intro"><p>{e(intro)}</p></div>' if intro else "")
        + f'<div class="content tight"><ul class="ruled compact qs-scope">{"".join(rows)}</ul>'
        + (f'<p class="qs-capacity">{e(capacity)}</p>' if capacity else "")
        + f'<div class="qs-price"><p class="incl-label">{e(L("price"))}</p>'
          f'<div class="amount">{e(amount)}</div><p class="basis">{e(basis)}</p>'
        + (f'<p class="validity">{e(validity)}</p>' if validity else "")
        + "</div></div>")


def s_work(s, w, kind):
    if kind != "curation":
        warn(w, 'a "work" slide names a film: only curation decks may ("kind": "curation")')
    facts = items(s, "facts", w)
    facts_html = ('<ul class="ruled terms facts">' + "".join(
        f'<li><span class="term">{e(f.get("term", ""))}</span><span>{e(f.get("text", ""))}</span></li>'
        for f in facts) + "</ul>") if facts else ""
    syn = field(s, "synopsis", w)
    fit = field(s, "fit", w)
    pillar = field(s, "pillar", w)
    ov = field(s, "overline", w, "eyebrow") or pillar
    link = s.get("link") or {}
    link_html = ""
    if link.get("href") and (DRAFT or not has_ph(link)):
        link_html = (f'<p class="work-link"><a href="{html_lib.escape(str(link["href"]))}">'
                     f'{e(link.get("label") or link["href"])}</a></p>')
    info = (f'<div class="info">' + (f'<div class="overline">{e(ov)}</div>' if ov else "")
            + heading(s, w)
            + (f'<p class="synopsis">{e(syn)}</p>' if syn else "")
            + (f'<p class="fit">{e(fit)}</p>' if fit else "")
            + link_html
            + f"{facts_html}</div>")
    credit = field(s, "image_credit", w)
    if s.get("image") and not credit:
        warn(w, "a film still needs an image_credit (where it comes from, e.g. the distributor's press kit)")
    media = (f'<div class="still">{aperture(s.get("image"), w, focal=s.get("focal"))}'
             + (f'<p class="t-cap">{e(credit)}</p>' if credit else "") + "</div>")
    if not s.get("sources") or has_ph(s.get("sources")):
        warn(w, "facts on this title card have no checked sources yet: run the audit (audit-deck.py)")
    return "work-slide", f'<div class="work">{media}{info}</div>'


def s_acceptance(s, w):
    parties = s.get("parties") or L("parties")
    text = field(s, "text", w)
    sig = L("sig")
    ps = "".join(f'<div class="party"><div class="label">{e(p)}</div>'
                 f'<div class="sig">{e(sig[0])}</div><div class="sig">{e(sig[1])}</div>'
                 f'<div class="sig" style="margin-bottom:0">{e(sig[2])}</div></div>' for p in parties)
    return "acceptance", (f'{overline(s, w)}{heading(s, w)}'
                          + (f'<p class="lead">{e(text)}</p>' if text else "")
                          + f'<div class="content tight"><div class="parties">{ps}</div></div>')


def s_closing(s, w):
    """One path only: write to hello@poeticsvr.com. No buttons, no second option."""
    title = field(s, "title", w) or L("closing")
    size = step_down(title, ["l", "m"], [24, 999])
    for k in ("cta", "secondary", "contact"):
        if s.get(k):
            warn(w, f'"{k}" ignored: the closing page offers one path, the email')
    mail = SITE["email"]
    return "closing", (f'<div class="band">{ring_svg()}<h1 class="h h-{size}">{e(title)}</h1>'
                       f'<a class="email" href="mailto:{mail}">{mail}</a></div>')


def s_text(s, w):
    paras = items(s, "paragraphs", w) or ([s["text"]] if isinstance(s.get("text"), str) else [])
    body = "".join(f"<p>{e(x)}</p>" for x in paras)
    return "text-slide", (f'{overline(s, w)}{heading(s, w)}{lead(s, w)}'
                          f'<div class="content"><div class="prose">{body}</div></div>')


RENDERERS = {
    "cover": s_cover, "title": s_cover, "statement": s_statement, "section": s_section,
    "agenda": s_agenda, "terms": s_terms, "columns": s_columns, "cards": s_columns,
    "stats": s_stats, "split": s_split, "timeline": s_timeline, "quote": s_quote,
    "people": s_people, "gallery": s_gallery, "media": s_media, "table": s_table,
    "price": s_price, "acceptance": s_acceptance, "closing": s_closing, "text": s_text,
    "quote_sheet": s_quote_sheet,
}


def dispatch(s, w, deck, density):
    stype = s.get("type", "content")
    if stype == "content":
        return s_content(s, w, density)
    if stype == "work":
        return s_work(s, w, deck.get("kind", "standard"))
    if stype in RENDERERS:
        return RENDERERS[stype](s, w)
    warn(w, f'unknown slide type "{stype}", rendered as content')
    return s_content(s, w, density)


def created_date(deck):
    import datetime
    if deck.get("date"):
        return str(deck["date"])
    d = datetime.date.today()
    if LANG == "el":
        return d.strftime("%d/%m/%Y")
    return f"{d.day} {d.strftime('%B %Y')}"


def render_slide(s, k, total, deck):
    s = resolve_template(s, f"slide {k}")
    stype = s.get("type", "content")
    w = f"slide {k} ({template_label(s)})"
    for text in walk_strings({kk: vv for kk, vv in s.items() if kk != "notes"}):
        lint_text(w, text)
    density = s.get("density") or deck.get("density", "speaker")
    lint_rules(w, s, deck, density)
    if stype == "page":
        out = []
        for j, sec in enumerate(s.get("sections") or []):
            sec = resolve_template(sec, f"{w} section {j + 1}")
            sw = f"{w} section {j + 1} ({template_label(sec)})"
            for text in walk_strings({kk: vv for kk, vv in sec.items() if kk != "notes"}):
                lint_text(sw, text)
            lint_rules(sw, sec, deck, density)
            scls, sbody = dispatch(sec, sw, deck, density)
            out.append(f'<div class="sec {scls.split()[0]}">{sbody}</div>')
        if FORMAT != "a4":
            warn(w, "stacked pages (A1) are made for A4; on 16:9 they will crowd")
        cls, body = "page-slide", f'<div class="flow">{"".join(out)}</div>'
    else:
        cls, body = dispatch(s, w, deck, density)

    if s.get("density") in ("reading", "speaker"):
        cls += f' {s["density"]}'
    ground = s.get("ground")
    if ground == "deep" and "deep" not in cls:
        cls += " deep"
    elif ground == "paper":
        cls = cls.replace(" deep", "")

    note = field(s, "note", w)
    small = f'<p class="small-print">{e(note)}</p>' if note else ""
    left = deck.get("footer", "poeticsvr.com")
    if deck.get("confidential", True):
        left = f"{left} · {L('confidential')}"
    proposal = deck.get("kind") == "proposal"
    if proposal or deck.get("footer_date"):
        left = f"{left} · {L('created')} {created_date(deck)}"
    right = [f'Nº {deck["number"]}'] if deck.get("number") else []
    if stype not in ("cover", "title", "closing") and total > 1:
        right.append(f"{k} {L('of')} {total}")
    disclaimer = deck.get("disclaimer", L("disclaimer") if proposal else "")
    disc = f'<span class="disclaimer">{e(disclaimer)}</span>' if disclaimer else ""
    if disclaimer:
        cls += " has-disclaimer"
    foot = f'<div class="foot">{disc}<span>{e(left)}</span><span>{e(" · ".join(right))}</span></div>'
    notes = s.get("notes")
    notes_html = ""
    if notes:
        paras = notes if isinstance(notes, list) else [notes]
        notes_html = '<aside class="notes">' + "".join(f"<p>{e(p)}</p>" for p in paras) + "</aside>"
    safe = 180 if disclaimer else 132
    return (f'<section class="slide {cls}" data-type="{e(stype)}" data-safe="{safe}">{wordmark_svg()}'
            f'{body}{small}{foot}{notes_html}</section>')


def lint_deck(deck):
    slides = [resolve_template(s, "deck") if s.get("template") in PAGE_TEMPLATES else s
              for s in deck.get("slides", [])]
    run = 1
    for k in range(1, len(slides)):
        same = slides[k].get("type") == slides[k - 1].get("type") and slides[k].get("type") in (
            "content", "columns", "cards", "terms", "split", "stats")
        run = run + 1 if same else 1
        if run == 3:
            warn(f"slide {k + 1}", f'third "{slides[k].get("type")}" slide in a row: vary the rhythm')
    flat = slides + [resolve_template(sec, "deck") for s in slides for sec in (s.get("sections") or [])]
    if deck.get("kind") == "proposal" and not any(s.get("type") in ("price", "table", "quote_sheet") for s in flat):
        if deck.get("mode") in ("offer", "tiered", "quote"):
            warn("deck", f'{deck.get("mode")} proposal without a price or table slide')
    if deck.get("kind") == "proposal":
        if any(s.get("type") in ("cover", "title", "closing") for s in slides):
            warn("deck", "proposals have no cover or closing page: start with the quote sheet (D3)")
    elif slides and slides[-1].get("type") != "closing":
        warn("deck", "the last slide isn't a closing slide")


def deck_format(deck):
    fmt = str(deck.get("format") or ("a4" if deck.get("kind") == "proposal" else "16:9")).lower()
    if fmt not in ("16:9", "a4"):
        warn("deck", f'unknown format "{fmt}", using 16:9')
        fmt = "16:9"
    return fmt


def build_html(deck: dict) -> str:
    global DRAFT, FORMAT, LANG
    DRAFT = bool(deck.get("draft"))
    FORMAT = deck_format(deck)
    LANG = deck.get("lang", "en") if deck.get("lang") in LABELS else "en"
    theme = deck.get("theme", "light")
    if theme not in THEMES:
        warn("deck", f'unknown theme "{theme}", using light')
        theme = "light"
    slides = deck.get("slides", [])
    lint_deck(deck)
    lint_rhythm(deck)
    total = len(slides)
    parts = [render_slide(s, k + 1, total, deck) for k, s in enumerate(slides)]
    css = CSS.replace("__LIGHT__", theme_vars("light")).replace("__DARK__", theme_vars("dark"))
    reading = " reading" if (deck.get("density") or ("reading" if FORMAT == "a4" else "")) == "reading" else ""
    reading += " a4" if FORMAT == "a4" else ""
    page_rule = ("@page { size: 210mm 297mm; margin: 0; }" if FORMAT == "a4"
                 else "@page { size: 1920px 1080px; margin: 0; }")
    title = e(deck.get("title", "Poetics"))
    return f"""<!DOCTYPE html>
<html lang="{e(deck.get("lang", "en"))}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<style>
{font_face_css()}
{css}
{page_rule}
</style>
</head>
<body>
<div class="deck{reading}" data-theme="{theme}">
<div class="stage">
{chr(10).join(parts)}
</div>
<div class="notes-panel" aria-live="polite"></div>
</div>
<script>{JS}</script>
</body>
</html>"""


# -- Browser (PDF, --check) ----------------------------------------------------------

def check_playwright():
    import importlib
    import os as _os
    import subprocess
    try:
        importlib.import_module("playwright")
    except ImportError:
        r = subprocess.run([sys.executable, "-m", "pip", "install", "--break-system-packages", "-q", "playwright"])
        if r.returncode != 0:
            subprocess.run([sys.executable, "-m", "pip", "install", "-q", "playwright"], check=True)
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as pw:
            exe = pw.chromium.executable_path
        if not (exe and _os.path.exists(exe)):
            raise RuntimeError
    except Exception:
        channel = _load_brand_assets().pick_browser_channel()
        if channel:
            # System Chrome/Edge renders identically: skip the download (which some
            # networks block outright).
            _os.environ.setdefault("POETICS_BROWSER_CHANNEL", channel)
            print(f"Bundled Chromium not installed: using system browser ({channel}).", file=sys.stderr)
        else:
            subprocess.run([sys.executable, "-m", "playwright", "install", "chromium"], check=True)


CHECK_JS = r"""
() => {
  const out = [];
  document.querySelectorAll('.slide').forEach((slide, k) => {
    const r = slide.getBoundingClientRect();
    const issues = new Set();
    slide.querySelectorAll('*').forEach(el => {
      if (el.closest('.foot, .wordmark, .ring, .notes, .small-print') || el.matches('svg, path, img')) return;
      if (el.closest('.bleed-right, .bleed-left, .cover-media')) return;
      const b = el.getBoundingClientRect();
      if (!b.width || !b.height) return;
      const label = (el.textContent || '').trim().replace(/\s+/g, ' ').slice(0, 50) || el.className;
      if (b.right > r.right - 120 + 2) issues.add('runs past the right margin: ' + label);
      if (b.left < r.left + 120 - 2) issues.add('runs past the left margin: ' + label);
      if (b.bottom > r.bottom - (+slide.dataset.safe || 132)) issues.add('crowds the footer: ' + label);
      if (b.bottom > r.bottom) issues.add('falls off the slide: ' + label);
    });
    // Guardrail: room around the wordmark. Nothing (text or photo) within 72px below it.
    const mark = slide.querySelector('.wordmark');
    if (mark) {
      const limit = mark.getBoundingClientRect().bottom + 72;
      slide.querySelectorAll('*').forEach(el => {
        if (el.closest('.wordmark, .foot, .notes, .ring') || el.matches('svg, path, img')) return;
        const b = el.getBoundingClientRect();
        const hasText = Array.from(el.childNodes).some(n => n.nodeType === 3 && n.textContent.trim());
        if (b.width && b.height && (hasText || el.classList.contains('ap')) && b.top < limit)
          issues.add('too close to the wordmark (' + Math.round(b.top - limit + 72) + 'px, needs 72): ' +
                     ((el.textContent || '').trim().slice(0, 40) || 'photo'));
      });
    }
    // Guardrail: nothing overlaps. Text against text, text against a photo.
    const leaves = Array.from(slide.querySelectorAll('*')).filter(el => {
      if (el.closest('.notes, .ring') || el.matches('svg, path, img')) return false;
      if (getComputedStyle(el).display === 'inline') return false;   // wrapped inline boxes span lines
      const t = Array.from(el.childNodes).some(n => n.nodeType === 3 && n.textContent.trim());
      return t || el.classList.contains('ap');
    }).map(el => [el, el.getBoundingClientRect()]).filter(([el, b]) => b.width && b.height);
    for (let i = 0; i < leaves.length; i++) for (let j = i + 1; j < leaves.length; j++) {
      const [a, ra] = leaves[i], [c, rc] = leaves[j];
      if (a.contains(c) || c.contains(a)) continue;
      if (a.classList.contains('ap') && c.classList.contains('ap')) continue;
      const ox = Math.min(ra.right, rc.right) - Math.max(ra.left, rc.left);
      const oy = Math.min(ra.bottom, rc.bottom) - Math.max(ra.top, rc.top);
      if (ox > 4 && oy > 4)
        issues.add('overlap: ' + ((a.textContent || '').trim().slice(0, 25) || 'photo') + ' / ' +
                   ((c.textContent || '').trim().slice(0, 25) || 'photo'));
    }
    const small = slide.querySelector('.small-print');
    if (small) {
      const sb = small.getBoundingClientRect();
      slide.querySelectorAll('.content *, .text *').forEach(el => {
        const b = el.getBoundingClientRect();
        if (b.height && b.bottom > sb.top + 2 && b.top < sb.bottom &&
            Array.from(el.childNodes).some(n => n.nodeType === 3 && n.textContent.trim()))
          issues.add('collides with the small print: ' + el.textContent.trim().slice(0, 40));
      });
    }
    if (issues.size) out.push([k + 1, Array.from(issues).slice(0, 4)]);
  });
  return out;
}
"""


def render_with_browser(html: str, *, pdf_path=None, png_dir=None, fmt="16:9"):
    """Render once in Chromium: optionally a PDF and/or one PNG per slide. Returns
    layout issues found by CHECK_JS (only when png_dir is given)."""
    import os as _os
    import tempfile
    from playwright.sync_api import sync_playwright
    issues = []
    with tempfile.TemporaryDirectory() as tmp:
        page_path = pathlib.Path(tmp) / "deck.html"
        page_path.write_text(html, encoding="utf-8")
        with sync_playwright() as p:
            browser = p.chromium.launch(channel=_os.environ.get("POETICS_BROWSER_CHANNEL") or None)
            page = browser.new_page(viewport={"width": 1920, "height": 1080})
            page.goto(page_path.as_uri() + "?static")
            page.evaluate("document.fonts.ready")
            page.wait_for_timeout(300)
            if pdf_path:
                if fmt == "a4":   # 1480px wide pages scaled onto real 210 x 297 mm sheets
                    page.pdf(path=str(pdf_path), width="210mm", height="297mm", scale=793.7 / 1480,
                             print_background=True, margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
                else:
                    page.pdf(path=str(pdf_path), width="1920px", height="1080px", print_background=True,
                             margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
            if png_dir:
                png_dir.mkdir(parents=True, exist_ok=True)
                for old in png_dir.glob("slide-*.png"):
                    old.unlink()
                issues = page.evaluate(CHECK_JS)
                for k, el in enumerate(page.query_selector_all(".slide")):
                    el.screenshot(path=str(png_dir / f"slide-{k + 1:02d}.png"))
            browser.close()
    return issues


# -- PPTX (python-pptx) ---------------------------------------------------------------

def export_pptx(deck: dict, out_path: str):
    """Editable approximation. Coordinates are slide px (1920 x 1080); 144 px = 1 inch,
    so 2 px = 1 pt."""
    import importlib
    import io
    import subprocess
    try:
        importlib.import_module("pptx")
    except ImportError:
        r = subprocess.run([sys.executable, "-m", "pip", "install", "python-pptx", "--break-system-packages", "-q"])
        if r.returncode != 0:
            subprocess.run([sys.executable, "-m", "pip", "install", "python-pptx", "-q"], check=True)
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Emu, Pt

    global DRAFT
    DRAFT = bool(deck.get("draft"))
    if deck_format(deck) == "a4":
        sys.exit("error: A4 decks (proposals) export as .html or .pdf, not .pptx")
    T = THEMES.get(deck.get("theme", "light"), THEMES["light"])
    reading = deck.get("density") == "reading"
    TXT = 26 if reading else 32
    DISPLAY, SANS = "Instrument Serif", "Inter Tight"

    def px(v):
        return Emu(int(v * 6350))  # 914400 EMU per inch / 144 px per inch

    def rgb(h):
        h = h.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1920), px(1080)
    prs.core_properties.title = plain(deck.get("title", "Poetics"))
    blank = prs.slide_layouts[6]

    def text(sl, x, y, w, h, value, *, size=26, font=SANS, color="ink", bold=False,
             anchor=None, line=None):
        box = sl.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if anchor:
            tf.vertical_anchor = anchor
        paras = value if isinstance(value, list) else [value]
        for k, val in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.text = plain(val)
            p.font.name, p.font.size, p.font.bold = font, Pt(size / 2), bold
            p.font.color.rgb = rgb(T[color])
            if line:
                p.line_spacing = line
            if k:
                p.space_before = Pt(size / 2 * 0.6)
        return box

    def rule(sl, x, y, w, color="shade"):
        ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x), px(y), px(x + w), px(y))
        ln.line.color.rgb = rgb(T[color])
        ln.line.width = Pt(0.75)

    def rect(sl, x, y, w, h, fill, radius=None):
        shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                  px(x), px(y), px(w), px(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(T[fill])
        shp.line.fill.background()
        shp.shadow.inherit = False
        if radius:
            shp.adjustments[0] = min(0.5, radius / min(w, h))
        return shp

    def picture(sl, src, x, y, w, h, radius=48, focal="center"):
        uri = resolve_image(src) if src else ""
        if not uri:
            rect(sl, x, y, w, h, "paper-deep", radius)
            return
        raw = base64.b64decode(uri.split(",", 1)[1])
        pic = sl.shapes.add_picture(io.BytesIO(raw), px(x), px(y), px(w), px(h))
        try:
            from PIL import Image
            iw, ih = Image.open(io.BytesIO(raw)).size
            target, actual = w / h, iw / ih
            fx, fy = [float(v.rstrip("%")) / 100 for v in FOCAL.get(focal or "center", "50% 50%").split()]
            if actual > target:
                excess = 1 - target / actual
                pic.crop_left, pic.crop_right = excess * fx, excess * (1 - fx)
            else:
                excess = 1 - actual / target
                pic.crop_top, pic.crop_bottom = excess * fy, excess * (1 - fy)
        except Exception:
            pass
        try:
            pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
            pic._element.spPr.prstGeom.avLst  # noqa: B018 - geometry exists
            adj = min(0.5, radius / min(w, h))
            from pptx.oxml.ns import qn
            av = pic._element.spPr.prstGeom.find(qn("a:avLst"))
            gd = av.makeelement(qn("a:gd"), {"name": "adj", "fmla": f"val {int(adj * 100000)}"})
            av.append(gd)
        except Exception:
            pass

    def header(sl, s, y=192, w=1680):
        ov = s.get("overline") or s.get("eyebrow")
        if ov and not (has_ph(ov) and not DRAFT):
            text(sl, 120, y - 56, w, 32, ov, size=21, color="ink-soft")
        title = s.get("title", "")
        if title and not (has_ph(title) and not DRAFT):
            lines = 1 if len(plain(title)) <= 34 else 2
            text(sl, 120, y, w, 80 * lines, title, size=66, font=DISPLAY, line=1.05)
            y += 76 * lines
        ld = s.get("lead") or s.get("subtitle")
        if ld and not (has_ph(ld) and not DRAFT):
            text(sl, 120, y + 40, min(w, 1100), 100, ld, size=32)
            y += 140
        return y + 80

    def keep(lst):
        return [v for v in (lst or []) if DRAFT or not has_ph(v)]

    slides = [resolve_template(s, "pptx") for s in deck.get("slides", [])]
    total = len(slides)
    for k, s in enumerate(slides, 1):
        stype = s.get("type", "content")
        sl = prs.slides.add_slide(blank)
        deep = stype in ("statement", "section", "quote") or s.get("ground") == "deep"
        if s.get("ground") == "paper":
            deep = False
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = rgb(T["paper-deep" if deep else "paper"])
        text(sl, 120, 66, 400, 40, "POETICS", size=26, bold=True)  # wordmark stand-in

        if stype in ("cover", "title"):
            img = s.get("image")
            tw = 960 if img else 1680
            if img:
                picture(sl, img, 1116, 120, 804, 840, focal=s.get("focal"))
            if s.get("overline"):
                text(sl, 120, 520, tw, 32, s["overline"], size=21, color="ink-soft")
            text(sl, 120, 570, tw, 300, s.get("title", ""), size=102 if img else 156, font=DISPLAY,
                 anchor=MSO_ANCHOR.BOTTOM, line=1.0)
            if s.get("lead") or s.get("subtitle"):
                text(sl, 120, 900, tw, 80, s.get("lead") or s.get("subtitle"), size=32, color="ink-soft")
        elif stype == "statement":
            if s.get("overline"):
                text(sl, 120, 300, 1680, 32, s["overline"], size=21, color="ink-soft")
            text(sl, 120, 350, 1500, 360, s.get("text") or s.get("title", ""), size=102, font=DISPLAY, line=1.05)
            facts = keep(s.get("facts"))
            if facts:
                text(sl, 120, 740, 1680, 50, "      ".join(plain(f) for f in facts), size=26, color="ink-soft")
            if s.get("credit"):
                text(sl, 120, 820, 1680, 50, s["credit"], size=26, color="ink-soft")
        elif stype == "section":
            if s.get("overline"):
                text(sl, 120, 640, 1680, 32, s["overline"], size=21, color="ink-soft")
            text(sl, 120, 690, 1680, 240, s.get("title", ""), size=102, font=DISPLAY, line=1.05,
                 anchor=MSO_ANCHOR.BOTTOM)
        elif stype == "quote":
            text(sl, 120, 260, 1500, 420, f'“{plain(s.get("quote", ""))}”', size=66, font=DISPLAY, line=1.15)
            who = [v for v in (s.get("name") or s.get("attribution"), s.get("role")) if v]
            if who:
                text(sl, 120, 760, 1200, 90, who, size=26)
        elif stype == "closing":
            rect(sl, 120, 300, 1680, 480, "paper-deep", radius=240)
            text(sl, 312, 380, 1300, 140, s.get("title") or "Write to us.", size=102, font=DISPLAY)
            mail = text(sl, 312, 560, 1300, 90, SITE["email"], size=66, font=DISPLAY, color="lens")
            mail.text_frame.paragraphs[0].runs[0].hyperlink.address = f"mailto:{SITE['email']}"
        else:
            if stype in ("split", "work"):
                img_x = 996 if s.get("side") != "left" else 120
                text_x_shift = 0 if s.get("side") != "left" else 876
                y = 192
                if stype == "work":
                    picture(sl, s.get("image"), 120, 168, 1000, 744, focal=s.get("focal"))
                    x0, wd = 1192, 608
                    if s.get("pillar"):
                        text(sl, x0, 168, wd, 32, s["pillar"], size=21, color="ink-soft")
                    text(sl, x0, 210, wd, 160, s.get("title", ""), size=66, font=DISPLAY, line=1.05)
                    body = [v for v in (s.get("synopsis"), s.get("fit")) if v]
                    text(sl, x0, 400, wd, 280, body, size=26)
                    fy = 700
                    for f in keep(s.get("facts"))[:5]:
                        text(sl, x0, fy, 180, 30, f.get("term", ""), size=21, color="ink-soft")
                        text(sl, x0 + 190, fy, wd - 190, 30, f.get("text", ""), size=21)
                        fy += 42
                else:
                    picture(sl, s.get("image"), img_x, 192, 804, 720, focal=s.get("focal"))
                    if s.get("title"):
                        text(sl, 120 + text_x_shift, 192, 804, 170, s["title"], size=66, font=DISPLAY, line=1.05)
                    y = 400
                    for it in keep(s.get("items")):
                        rule(sl, 120 + text_x_shift, y, 804)
                        text(sl, 120 + text_x_shift, y + 24, 804, 50, it.get("title", ""), size=TXT, bold=True)
                        if it.get("text") or it.get("body"):
                            text(sl, 120 + text_x_shift, y + 24 + TXT * 1.6, 804, 60, it.get("text") or it.get("body"),
                                 size=26, color="ink-soft")
                        y += 130
            else:
                y = header(sl, s)
                if stype in ("content", "agenda", "terms"):
                    rows = keep(s.get("bullets") or s.get("items") or s.get("rows"))
                    for n, r in enumerate(rows):
                        rule(sl, 120, y, 1680)
                        if stype == "terms":
                            text(sl, 120, y + 30, 420, 60, r.get("term", ""), size=TXT, bold=True)
                            text(sl, 576, y + 30, 1224, 60, r.get("text", ""), size=TXT, color="ink-soft")
                        elif stype == "agenda":
                            text(sl, 120, y + 36, 72, 40, f"{n + 1:02d}", size=21, color="ink-soft")
                            text(sl, 192, y + 30, 1608, 60, r, size=TXT)
                        else:
                            text(sl, 120, y + 30, 1680, 60, r, size=TXT)
                        y += 60 + TXT * 1.5
                    rule(sl, 120, y, 1680)
                elif stype in ("columns", "cards", "stats", "timeline", "people"):
                    key = {"columns": "columns", "cards": "cards", "stats": "stats",
                           "timeline": "steps", "people": "people"}[stype]
                    cols = keep(s.get(key) or s.get("cards"))
                    n = max(1, len(cols) if len(cols) <= 4 else 3)
                    cw = (1680 - 36 * (n - 1)) / n
                    for j, c in enumerate(cols):
                        cx = 120 + (j % n) * (cw + 36)
                        cy = y + (j // n) * 300
                        if stype == "people":
                            picture(sl, c.get("image"), cx, cy, cw, 912 - cy - 110, focal=c.get("focal") or "top")
                            text(sl, cx, 912 - 96, cw, 100,
                                 [v for v in (c.get("name"), c.get("role"), c.get("context")) if v], size=21)
                            continue
                        rule(sl, cx, cy, cw, "ink" if c.get("now") else "shade")
                        if stype == "stats":
                            text(sl, cx, cy + 30, cw, 150, c.get("value", ""), size=144 if len(cols) <= 3 else 102,
                                 font=DISPLAY, line=1.0)
                            text(sl, cx, cy + 190, cw, 80, [v for v in (c.get("label"), c.get("sublabel")) if v],
                                 size=26, color="ink-soft")
                        else:
                            if stype == "timeline":
                                text(sl, cx, cy + 30, cw, 30, "Now" if c.get("now") else f"{j + 1:02d}", size=21,
                                     color="ink" if c.get("now") else "ink-soft")
                                cy += 48
                            text(sl, cx, cy + 30, cw, 50, c.get("title", ""), size=36, bold=True)
                            body = [v for v in (c.get("text") or c.get("body"), c.get("when")) if v]
                            if body:
                                text(sl, cx, cy + 100, cw, 200, body, size=26, color="ink-soft")
                elif stype in ("gallery", "media"):
                    imgs = s.get("images") or [s.get("image")]
                    imgs = [im if isinstance(im, str) else (im or {}).get("src") for im in imgs][:8]
                    n = len(imgs)
                    cols, rows = {1: (1, 1), 2: (2, 1), 3: (3, 1), 4: (2, 2), 5: (3, 2), 6: (3, 2)}.get(n, (4, 2))
                    top = y if s.get("title") else 192
                    gw, gh = (1680 - 36 * (cols - 1)) / cols, (912 - top - 36 * (rows - 1)) / rows
                    for j, im in enumerate(imgs):
                        picture(sl, im, 120 + (j % cols) * (gw + 36), top + (j // cols) * (gh + 36), gw, gh)
                elif stype == "table":
                    cols = s.get("columns") or []
                    rows = keep(s.get("rows"))
                    ncol = max([len(cols)] + [len(r) for r in rows] + [1])
                    shape = sl.shapes.add_table(len(rows) + (1 if cols else 0), ncol, px(120), px(y),
                                                px(1680), px(70 * (len(rows) + 1)))
                    tbl = shape.table
                    all_rows = ([cols] if cols else []) + rows
                    for r_i, row in enumerate(all_rows):
                        for c_i in range(ncol):
                            cell = tbl.cell(r_i, c_i)
                            cell.text = plain(row[c_i]) if c_i < len(row) else ""
                            cell.fill.solid()
                            hl = s.get("highlight") is not None and c_i == s.get("highlight")
                            cell.fill.fore_color.rgb = rgb(T["paper-deep" if hl else ("paper-deep" if deep else "paper")])
                            p = cell.text_frame.paragraphs[0]
                            p.font.name, p.font.size = SANS, Pt(10.5 if (cols and r_i == 0) else 13)
                            p.font.color.rgb = rgb(T["ink-soft" if (cols and r_i == 0) else "ink"])
                            p.font.bold = c_i == 0 and not (cols and r_i == 0)
                elif stype == "price":
                    text(sl, 120, y, 1000, 160, s.get("amount", ""), size=144, font=DISPLAY)
                    text(sl, 120, y + 180, 1000, 50, s.get("basis") or "All-inclusive, excl. VAT", size=32, color="ink-soft")
                    if s.get("validity"):
                        text(sl, 120, y + 260, 900, 90, s["validity"], size=26, color="ink-soft")
                    iy = y
                    for it in keep(s.get("includes")):
                        rule(sl, 1176, iy, 624)
                        text(sl, 1176, iy + 18, 624, 40, it, size=26)
                        iy += 78
                elif stype == "acceptance":
                    for j, p in enumerate(s.get("parties") or ["For the venue", "For Poetics"]):
                        cx = 120 + j * 888
                        text(sl, cx, y + 40, 792, 40, p, size=26, bold=True)
                        for m, lbl in enumerate(("Name and role", "Signature", "Date")):
                            rule(sl, cx, y + 200 + m * 130, 792, "ink-soft")
                            text(sl, cx, y + 212 + m * 130, 792, 30, lbl, size=21, color="ink-soft")

        left = deck.get("footer", "poeticsvr.com") + (" · Confidential" if deck.get("confidential", True) else "")
        right = " · ".join(([f'Nº {deck["number"]}'] if deck.get("number") else [])
                           + ([f"{k} of {total}"] if stype not in ("cover", "title", "closing") else []))
        text(sl, 120, 1010, 900, 30, left, size=21, color="ink-soft")
        if right:
            text(sl, 1000, 1010, 800, 30, right, size=21, color="ink-soft").text_frame.paragraphs[0].alignment = 3
        if s.get("note"):
            text(sl, 120, 950, 1200, 40, s["note"], size=21, color="ink-soft")
        notes = s.get("notes")
        if notes:
            sl.notes_slide.notes_text_frame.text = "\n\n".join(notes if isinstance(notes, list) else [notes])

    prs.save(out_path)


# -- Contact sheet ----------------------------------------------------------------------

def contact_sheet(png_dir, slides, out_png, cols=3):
    """Every slide on one page, captioned with its number and page template."""
    from PIL import Image, ImageDraw, ImageFont
    files = sorted(png_dir.glob("slide-*.png"))
    if not files:
        return
    w, h, gap, cap = 640, 360, 24, 40
    first = Image.open(files[0])
    if first.height > first.width:   # A4 pages: portrait thumbnails, four across
        w, h, cols = 360, int(360 * first.height / first.width), 4
    rows = (len(files) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * w + (cols + 1) * gap, rows * (h + cap) + (rows + 1) * gap),
                      THEMES["light"]["paper-deep"])
    draw = ImageDraw.Draw(sheet)
    try:
        font = ImageFont.truetype(str(FONTS / "InterTight-Variable.ttf"), 18)
    except Exception:
        font = ImageFont.load_default()
    for i, f in enumerate(files):
        x = gap + (i % cols) * (w + gap)
        y = gap + (i // cols) * (h + cap + gap)
        sheet.paste(Image.open(f).convert("RGB").resize((w, h), Image.LANCZOS), (x, y))
        label = template_label(slides[i]) if i < len(slides) else ""
        draw.text((x, y + h + 10), f"{i + 1:02d}  {label}", fill=THEMES["light"]["ink-soft"], font=font)
    sheet.save(out_png)


# -- Entry point ------------------------------------------------------------------------

def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = {a for a in sys.argv[1:] if a.startswith("--")}
    if len(args) != 2:
        sys.exit("usage: build-deck.py <deck.json> <output.html|output.pdf|output.pptx> [--check] [--sheet]")
    deck_path, out_path = pathlib.Path(args[0]), pathlib.Path(args[1])
    deck = json.loads(deck_path.read_text(encoding="utf-8"))
    _ALLOWED_EXTRA.append(deck_path.resolve().parent)

    ext = out_path.suffix.lower()
    if ext not in (".html", ".pdf", ".pptx"):
        sys.exit(f"error: unsupported output extension '{ext}' (use .html, .pdf, or .pptx)")

    html = build_html(deck)
    issues = []
    want_png = "--check" in flags or "--sheet" in flags
    png_dir = out_path.with_name(out_path.stem + "-check") if want_png else None

    if ext == ".html":
        out_path.write_text(html, encoding="utf-8")
    elif ext == ".pptx":
        export_pptx(deck, str(out_path))
    if ext == ".pdf" or png_dir:
        check_playwright()
        issues = render_with_browser(html, pdf_path=out_path if ext == ".pdf" else None, png_dir=png_dir,
                                     fmt=FORMAT)
    print(f"Wrote {out_path}")
    if png_dir:
        print(f"Slide images for review: {png_dir}/")
    if "--sheet" in flags and png_dir:
        sheet = out_path.with_name(out_path.stem + "-sheet.png")
        contact_sheet(png_dir, [resolve_template(s, "sheet") for s in deck.get("slides", [])], sheet)
        print(f"Contact sheet: {sheet}")

    seen = []
    for w in WARNINGS:
        if w not in seen:
            seen.append(w)
    if seen:
        print(f"\n{len(seen)} warning(s):", file=sys.stderr)
        for w in seen:
            print(f"  - {w}", file=sys.stderr)
    if issues:
        print(f"\nLayout problems on {len(issues)} slide(s):", file=sys.stderr)
        for k, msgs in issues:
            for m in msgs:
                print(f"  - slide {k}: {m}", file=sys.stderr)
    if png_dir and not issues:
        print("Layout check: nothing overflows.")
    fill = [(k + 1, filler_code(sl)) for k, sl in enumerate(deck.get("slides", []))]
    fill = [(k, c) for k, c in fill if c]
    if fill:
        print("Filler pages (ask the user whether to keep them): "
              + ", ".join(f"{k} {c} {PAGE_TEMPLATES[c][1]}" for k, c in fill))


if __name__ == "__main__":
    main()
